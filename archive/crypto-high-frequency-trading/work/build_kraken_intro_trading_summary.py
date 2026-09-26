from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/yitwah/Documents/Codex/2026-06-07/video-youtube-video-open-document-recently/outputs/kraken_introduction_to_crypto_trading_summary_zh.docx")


def rgb(value):
    return RGBColor.from_string(value)


def set_font(run, size=10, color="000000", bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def cell_text(cell, text, bold=False, size=8.1, color="000000", fill=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            set_cell_width(row.cells[i], width)


def add_table(doc, headers, rows, widths, size=7.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell_text(table.rows[0].cells[i], header, bold=True, size=8.3, color="0B2545", fill="E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, size=size)
    set_table_widths(table, widths)
    doc.add_paragraph("")
    return table


def add_callout(doc, title, body, fill="FFF4D6", color="5C3B00"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(f"{title}：")
    set_font(r, size=9.2, color=color, bold=True)
    r2 = p.add_run(body)
    set_font(r2, size=9.2, color=color)
    doc.add_paragraph("")


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=10)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=9.6)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=9.6)


def source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{label}: {url}")
    set_font(r, size=8.2, color="555555")


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, bold, before, after in [
        ("Title", 23, "0B2545", False, 0, 8),
        ("Heading 1", 15.5, "0B2545", True, 13, 6),
        ("Heading 2", 12.5, "1F4D78", True, 10, 5),
        ("Heading 3", 11, "1F4D78", True, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build():
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Kraken《Introduction to cryptocurrency trading》中文整理")

    sub = doc.add_paragraph()
    r = sub.add_run("官方更新时间：2026-03-09｜整理日期：2026-06-07｜用途：新手第一次理解交易所、订单、费用和 API 入口")
    set_font(r, size=9.5, color="555555")

    add_callout(
        doc,
        "一句话结论",
        "这篇 Kraken 文章不是深度交易教程，而是一张新手入口地图：先理解交易所能做什么，再按顺序学习币种、买卖流程、术语、订单类型、费用、杠杆/保证金和 API。它最适合放在学习路线的第一步。",
        "E8F4FD",
        "0B2545",
    )
    add_callout(
        doc,
        "风险提醒",
        "Kraken 官方也强调交易前要自己研究并理解风险。Crypto 价格波动大，交易可能亏损；部分产品和市场不受传统监管保护；税务、地区限制和产品可用性也会因司法辖区不同而变化。",
    )

    doc.add_heading("1. 原文核心内容", level=1)
    para(doc, "Kraken 把 cryptocurrency trading 定义为买入、卖出或持有 BTC、ETH 等数字资产，希望从市场价格变化中获利。它同时提到两类交易方式：用 USD、EUR、CAD 等现金换成 crypto，或者用一种 crypto 换另一种 crypto。")
    para(doc, "文章对 Kraken 的定位是全球加密货币交易所，强调安全记录、可交易数字资产范围和面向不同经验水平的交易工具：新手可以使用简单交易功能，经验更高的交易者可以使用 Kraken Pro、订单类型、API 等高级入口。")

    doc.add_heading("2. 新手应该按什么顺序读", level=1)
    add_table(
        doc,
        ["顺序", "主题", "要解决的问题", "新手读法"],
        [
            ("1", "BTC vs ETH / 主要资产", "先知道自己交易的资产是什么，不要只看涨跌。", "先读 BTC、ETH、稳定币和主流链，不急着买小币。"),
            ("2", "如何买 crypto", "理解入金、买入、卖出、提现和钱包的基本流程。", "先用小额现货，确认每一步资金路径。"),
            ("3", "交易术语表", "理解 pair、bid/ask、spread、volume、liquidity、maker/taker。", "看盘口前先把术语过一遍。"),
            ("4", "可交易币种列表", "知道交易所支持哪些资产和市场对。", "不要因为上架就认为安全或有投资价值。"),
            ("5", "杠杆和保证金", "理解借钱交易、爆仓和额外费用。", "新手先跳过，现货熟悉后再学。"),
            ("6", "交易费用", "知道成交才收费，费用受 30 日交易量、交易对和 maker/taker 影响。", "每笔交易前把手续费算进盈亏。"),
            ("7", "套利", "理解不同市场价格差可能产生机会。", "只当概念学习，不手动搬砖重仓。"),
            ("8", "订单类型和 API", "进阶到限价、止损、自动化和程序交易。", "先模拟，不给 API key 提现权限。"),
        ],
        [0.55, 1.35, 3.05, 2.55],
        7.8,
    )

    doc.add_heading("3. 交易所、订单和费用怎么理解", level=1)
    doc.add_heading("3.1 Kraken 作为交易所的作用", level=2)
    for item in [
        "撮合买方和卖方：你不是直接和 Kraken 对赌，常见交易是在订单簿中和其他交易者成交。",
        "支持现金/crypto 与 crypto/crypto 的兑换：例如 USD 买 BTC，或 BTC 换 ETH。",
        "提供不同入口：简单买卖适合新手，Kraken Pro、Desktop、API 更适合有经验的交易者。",
        "提供市场数据和账户工具：价格、交易对、订单、成交、余额、费用记录等。",
    ]:
        bullet(doc, item)

    doc.add_heading("3.2 订单类型入口", level=2)
    para(doc, "原文把“order types and options”作为进阶资源。Kraken 的订单类型页面列出 market、limit、stop loss、take profit、trailing stop、iceberg、conditional close 等。新手不要一开始全学，先掌握 market 和 limit，再学 stop loss。")
    add_table(
        doc,
        ["订单概念", "中文理解", "新手风险"],
        [
            ("Market order", "按市场当前可成交价格立即成交。", "成交快，但波动大或流动性差时滑点明显。"),
            ("Limit order", "只在指定价格或更好价格成交。", "可控价格，但可能不成交或只成交一部分。"),
            ("Stop loss", "价格触发后用于退出亏损头寸。", "触发价不等于最终成交价，剧烈行情仍会滑点。"),
            ("Take profit", "价格到达目标后锁定利润。", "过早止盈可能错过趋势，过晚可能利润回吐。"),
            ("Trailing stop", "止损价跟随行情移动。", "参数过近会被噪音扫出，过远又保护不足。"),
            ("Iceberg", "隐藏大订单的部分可见数量。", "新手一般不需要，适合大单执行。"),
        ],
        [1.25, 3.1, 3.45],
        7.7,
    )

    doc.add_heading("3.3 费用入口", level=2)
    para(doc, "Kraken 的费用说明强调：订单被执行、也就是与另一方订单匹配时才产生交易费；未成交就取消的订单通常不产生交易费。费用取决于 30 日交易量、交易对，以及订单是 maker 还是 taker。")
    for item in [
        "Maker：通常是挂单提供流动性，费用可能低一些。",
        "Taker：通常是吃掉已有订单，成交确定性更高，但费用可能更高。",
        "小额交易要注意最低交易费、价格精度和数量精度。",
        "保证金交易还会涉及开仓费、展期费等额外成本，新手先不要碰。",
    ]:
        bullet(doc, item)

    doc.add_heading("4. API 和自动化交易入口", level=1)
    para(doc, "原文把 Kraken API 作为经验交易者资源。Kraken 的 API 页面进一步说明，API 是把交易者的软件与 Kraken 的市场数据流或账户连接起来的接口；可以用于市场分析、自动下单和高频响应。")
    add_table(
        doc,
        ["API 类型/用途", "适合做什么", "新手注意"],
        [
            ("REST", "同步请求：查余额、下单、撤单、查订单、资金流程。", "适合先学；注意 rate limit 和错误重试。"),
            ("WebSockets", "实时数据流和异步交易信息。", "适合行情监控和更快响应；代码复杂度更高。"),
            ("FIX 4.4", "机构级低延迟交易接口。", "新手不需要，除非做机构系统。"),
            ("Bot / algorithmic trading", "用程序把市场数据转换成执行策略。", "必须人工监控，不能把 bot 当自动赚钱机器。"),
        ],
        [1.25, 3.3, 3.25],
        7.8,
    )
    add_callout(
        doc,
        "API 安全底线",
        "创建 API key 时只给必要权限；新手不要开提现权限；用独立子账户和小资金测试；所有订单、错误、余额变化都要记录；网络异常或订单状态未知时，程序应停机而不是继续加仓。",
        "FFF4D6",
        "5C3B00",
    )

    doc.add_heading("5. 读完后怎么行动", level=1)
    for step in [
        "先读 Kraken 文章和术语表，不做交易，只把 pair、spread、maker/taker、market/limit、liquidity 写进自己的词汇表。",
        "打开一个交易对页面，只观察盘口、成交、价格变化和费用，不下单。",
        "用小额现货学习一次完整流程：入金、买入、卖出、提现或留在交易所。",
        "只使用 market 和 limit 两类基础订单；能解释滑点和未成交后，再学 stop loss。",
        "把每笔交易前后的手续费、成交价、滑点、持仓理由写进日志。",
        "如果要用 API，先只读公开行情；之后再用只读 key；最后才用交易权限 key 做极小资金测试。",
    ]:
        number(doc, step)

    doc.add_heading("6. 这篇文章的局限", level=1)
    para(doc, "这篇文章更像 Kraken 学习中心的导航页，正文并不深入讲 K 线、技术分析、仓位管理、链上钱包安全或策略回测。因此它适合做第一站，但不能替代完整交易课程。")
    add_table(
        doc,
        ["缺口", "应该补什么"],
        [
            ("风险管理", "仓位、最大回撤、止损、交易日志、不要重仓单一资产。"),
            ("市场结构", "订单簿深度、流动性、价差、成交量真假、资金费率。"),
            ("安全", "2FA、白名单、API 权限、提现地址确认、防钓鱼。"),
            ("自动化", "回测、dry-run、异常处理、日志、停机机制。"),
        ],
        [1.6, 6.1],
        8.0,
    )

    doc.add_heading("7. 参考链接", level=1)
    source(doc, "Kraken: Introduction to cryptocurrency trading", "https://support.kraken.com/articles/360000674406-introduction-to-cryptocurrency-trading")
    source(doc, "Kraken: Order types & options", "https://support.kraken.com/sections/200577136-order-types")
    source(doc, "Kraken: How trading fees work", "https://support.kraken.com/articles/201893638-how-trading-fees-work-on-kraken")
    source(doc, "Kraken: API trading", "https://www.kraken.com/institutions/api")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
