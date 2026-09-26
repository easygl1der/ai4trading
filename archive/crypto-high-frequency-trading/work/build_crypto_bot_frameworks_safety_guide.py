from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path("/Users/yitwah/Documents/Codex/2026-06-07/video-youtube-video-open-document-recently")
OUT = BASE / "outputs/crypto_bot_frameworks_and_mcp_safety_guide_zh.docx"


def rgb(value):
    return RGBColor.from_string(value)


def set_font(run, size=10, color="111827", bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(int(w * 1440) for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    for row in table.rows:
        for i, width in enumerate(widths):
            set_cell_width(row.cells[i], width)
            set_cell_margins(row.cells[i])
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def cell_text(cell, text, *, bold=False, size=8.2, color="111827", fill=None, align=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align:
        p.alignment = align
    r = p.add_run(str(text))
    set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, *, size=8.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell_text(table.rows[0].cells[i], header, bold=True, size=8.4, color="0B2545", fill="E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, size=size)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def add_callout(doc, title, body, fill="E8F4FD", color="0B2545"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=170, end=170)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.14
    r = p.add_run(f"{title}：")
    set_font(r, size=9.3, color=color, bold=True)
    r2 = p.add_run(body)
    set_font(r2, size=9.3, color=color)
    set_table_geometry(table, [6.5])
    doc.add_paragraph("")


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.14
    r = p.add_run(text)
    set_font(r, size=10)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3.2)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=9.5)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3.2)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=9.5)


def source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(f"{label}: {url}")
    set_font(r, size=8.1, color="4B5563")


def setup_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    sec.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb("111827")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.14

    for name, size, color, bold, before, after in [
        ("Title", 23, "0B2545", False, 0, 8),
        ("Heading 1", 15.5, "0B2545", True, 13, 6),
        ("Heading 2", 12.5, "1F4D78", True, 10, 5),
        ("Heading 3", 11, "1F4D78", True, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Crypto bot framework safety guide | generated 2026-06-07")
    set_font(r, size=8, color="6B7280")


def build():
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Crypto 交易 Bot 框架、开源仓库与 MCP 安全选择指南")

    sub = doc.add_paragraph()
    r = sub.add_run("面向小白/初级程序员：哪些工具覆盖数据、信号、风控、执行、日志、监控；哪些能接 AI/MCP；哪些不该直接实盘")
    set_font(r, size=9.5, color="555555")

    add_callout(
        doc,
        "结论",
        "有现成框架能覆盖大部分模块，但没有任何框架能让交易天然安全。对你现在最稳的路线是：Freqtrade 做低频策略和 dry-run；Jesse 做回测研究和本地 MCP 辅助；Hummingbot 学做市/套利；CCXT 只当底层 API；NautilusTrader 留给更高级的事件驱动系统。AI/MCP 只能先做 read-only、paper trading 或本地研究，不应该直接拿交易权限控制真钱。",
    )

    add_callout(
        doc,
        "安全定义",
        "这里说的“安全”不是不会亏钱，而是工程上减少灾难：默认模拟、权限最小化、可回测、可停机、可审计、日志完整、不会把私钥/API key 暴露给 LLM 或陌生插件。",
        "FFF4D6",
        "5C3B00",
    )

    doc.add_heading("1. 你要的 6 层模块，成熟框架怎么覆盖", level=1)
    add_table(
        doc,
        ["层", "要解决的问题", "成熟框架通常怎么做", "安全检查"],
        [
            ("数据层", "行情、K 线、订单簿、成交、余额", "连接交易所 API，下载历史数据，实时订阅或轮询", "时间戳、缺失 K 线、API 限速、交易对精度"),
            ("信号层", "何时买/卖/观望", "策略类、指标、规则、模型或回调函数", "避免未来函数、过拟合、只在牛市有效"),
            ("风控层", "仓位、止损、最大亏损、冷却", "stoploss、stake sizing、protections、risk engine", "先能停，再能赚；默认小仓位"),
            ("执行层", "下单、撤单、订单状态", "统一订单接口、执行引擎、重试、订单生命周期", "半成交、重复下单、网络超时、交易所宕机"),
            ("记录层", "复盘、审计、找错误", "trades、logs、SQLite/Postgres、dashboard、导出", "每笔订单和错误都能追溯"),
            ("监控层", "异常发现和停机", "UI/API/Telegram/webhook、状态页、告警", "不能只会运行；必须能暂停/kill switch"),
        ],
        [0.7, 1.35, 2.35, 2.1],
        size=7.6,
    )

    doc.add_heading("2. 推荐框架总表", level=1)
    add_table(
        doc,
        ["工具", "最适合", "安全能力", "主要边界"],
        [
            ("Freqtrade", "低频/中频策略、现货/期货规则交易", "dry-run、backtesting、stoploss、protections、FreqUI/REST/Telegram", "不是专业做市或复杂低延迟系统"),
            ("Jesse", "策略研究、多品种/多周期回测、AI 辅助分析", "强调无 look-ahead bias，支持本地 MCP 辅助", "不适合新手直接做市或高频"),
            ("Hummingbot", "market making、grid、跨所/链上策略学习", "paper trade、kill switch、V2 controllers/executors、官方 MCP", "MCP 可触及实盘动作，新手不能直接放权"),
            ("NautilusTrader", "高级程序员、多资产、多场所严肃系统", "RiskEngine、ExecutionEngine、研究到实盘一致性设计", "学习成本高，不是第一套 bot"),
            ("CCXT", "自己写行情抓取、余额、下单、撤单", "内置 rate limiter，统一交易所 API", "不是完整 bot，需自己补风控/日志/监控"),
        ],
        [0.95, 1.95, 2.0, 1.6],
        size=7.45,
    )
    para(doc, "粗略排序：新手第一套优先 Freqtrade；研究和 AI 辅助可看 Jesse；做市/套利方向看 Hummingbot；自己裸写只把 CCXT 当底层；NautilusTrader 留给能维护事件驱动系统的人。")

    doc.add_heading("3. 最推荐：Freqtrade 作为第一套安全实验框架", level=1)
    para(doc, "如果你想把数据层、信号层、风控层、执行层、记录层、监控层放到一个框架里，Freqtrade 是最适合小白起步的选择。它不是保证盈利的工具，但工程结构完整，学习曲线比自己从 CCXT 裸写低很多。")
    for item in [
        "数据/策略：Python strategy class，处理 OHLCV、指标、entry/exit signal。",
        "回测/dry-run：先下载历史数据回测，再用 dry-run 接实时行情但不真实下单。",
        "风控：stoploss、ROI、stake size、cooldown、StoplossGuard 等 protections。",
        "执行/记录：处理交易生命周期，保存结果，可导出回测和成交记录。",
        "监控：FreqUI、REST API、Telegram、webhook；REST API 官方警告不要暴露到公网。",
    ]:
        bullet(doc, item)
    add_callout(
        doc,
        "Freqtrade 安全配置底线",
        "第一阶段只用 dry_run；第二阶段只给 API key 交易权限，不给提现权限；REST API 只监听 localhost；不要把 FreqUI 暴露到公网；任何策略先跑 lookahead/recursive 检查和至少两周 dry-run。",
        "FDEBD0",
        "6B3D00",
    )

    doc.add_heading("4. Jesse：适合研究、回测和本地 MCP 辅助", level=1)
    para(doc, "Jesse 的优势是研究和回测体验，它官方文档强调多周期/多交易对 backtest 和 livetrade，并明确提供 MCP 文档。Jesse 的 MCP 更适合让 AI 助手读取本地 Jesse 上下文、帮你改策略、跑回测、查设置，而不是直接让 AI 管真钱。")
    for item in [
        "适合：策略研究、图表、回测指标、参数优化、让 AI 帮忙理解策略文件。",
        "MCP 边界：Jesse MCP 服务在本机运行，需要你配置工具才会连接；你仍然决定何时开启。",
        "安全建议：MCP 初期只用于读取项目、跑回测、解释结果；不要一开始开放实盘交易动作。",
    ]:
        bullet(doc, item)

    doc.add_heading("5. Hummingbot：做市/套利/网格方向最专业", level=1)
    para(doc, "如果你关心做市、挂单、grid、cross-exchange market making，Hummingbot 比 Freqtrade 更贴近这个方向。它的 V2 controllers 会从 OrderBook、Trades、Candles 等市场数据生成 ExecutorActions，执行器负责下单和管理头寸。")
    for item in [
        "适合：market making、grid、跨所套利、DEX/connector 方向。",
        "安全能力：paper_trade 支持不用交易所 API 运行若干策略；kill switch 可以在达到正/负收益阈值时自动停止；适合先模拟库存和 spread。",
        "MCP 边界：Hummingbot 官方 MCP 可以让 AI 助手查看余额、订单、行情，也能下单、管理仓位、执行 swap、部署 bot；所以它不是“只读小插件”，新手必须先当高风险工具处理。",
        "风险：做市最大风险不是代码能不能跑，而是库存偏移、逆向选择、撤单延迟和手续费。",
    ]:
        bullet(doc, item)

    doc.add_heading("6. NautilusTrader：高级版事件驱动交易系统", level=1)
    para(doc, "NautilusTrader 是更专业的开源交易系统框架，官方定位是 Rust-native、Python 控制面的多资产多场所系统，覆盖研究、确定性仿真和 live execution。它有 RiskEngine、ExecutionEngine、ExecutionClient 等更严肃的工程组件。")
    add_table(
        doc,
        ["优点", "为什么重要"],
        [
            ("研究到实盘一致性", "同一事件驱动架构减少回测和实盘代码不一致。"),
            ("RiskEngine", "订单进入执行前可做 pre-trade risk checks。"),
            ("ExecutionEngine", "更清楚地管理订单生命周期和多场所执行。"),
            ("缺点", "学习成本高，不适合完全小白作为第一套 crypto bot。"),
        ],
        [1.45, 5.05],
        size=8.0,
    )

    doc.add_heading("7. CCXT：必要底层库，但不是安全框架", level=1)
    para(doc, "CCXT 很重要，但它不是完整 bot。它主要帮你统一交易所 API：行情、订单簿、余额、下单、撤单、市场规则等。官方 manual 提到有内置 rate limiter，但这不等于完整风控。")
    for item in [
        "适合：学交易所 API、写数据抓取、做低频执行层原型。",
        "缺少：策略生命周期、回测、风控、日志、监控、异常停机。",
        "安全建议：先 public API；再只读 key；最后交易 key；永远不开提现权限。",
    ]:
        bullet(doc, item)

    doc.add_heading("8. MCP 和 AI Agent：可以用，但要分级", level=1)
    para(doc, "MCP 的价值是让 AI 助手调用本地应用或工具，但交易场景尤其危险。交易 MCP 如果能下单、撤单、设置 API key，就等于把资金控制权交给一个工具链。新手应把 MCP 当成研究和辅助开发接口，不当自动交易入口。")
    add_table(
        doc,
        ["等级", "允许做什么", "例子", "风险判断"],
        [
            ("Level 0", "只读文档/代码/配置 schema", "Freqtrade introspection MCP、官方 docs 搜索", "推荐，适合写策略时查接口"),
            ("Level 1", "读取本地回测项目、运行 backtest", "Jesse 本地 MCP、只读本地工具", "可用，但要确认不会触发实盘"),
            ("Level 2", "连接 paper trading / sandbox", "Hummingbot paper trade、交易所 testnet", "可以实验，但仍要日志和限额"),
            ("Level 3", "真实账户下单/撤单", "Hummingbot MCP、交易所 trade MCP、AI agent 自动执行", "不建议新手；必须人工确认、限额和 kill switch"),
            ("禁止", "提现、导入私钥、无限授权、远程公网暴露", "陌生 Telegram bot、闭源 MCP、热钱包私钥", "直接拒绝"),
        ],
        [0.72, 1.75, 2.18, 1.85],
        size=7.8,
    )
    add_callout(
        doc,
        "MCP 安全底线",
        "优先使用官方或可审计开源 MCP；默认 read-only；不要让 LLM 看到交易所 secret；不要把 MCP server 暴露到公网；不要安装不明来源 MCP；任何 mutating tool 都要人工确认和金额上限。",
        "FDEBD0",
        "6B3D00",
    )

    doc.add_heading("9. 推荐组合", level=1)
    add_table(
        doc,
        ["你的目标", "推荐组合", "原因"],
        [
            ("小白第一套 bot", "Freqtrade + dry-run + FreqUI 本地", "覆盖最完整，默认能从回测到模拟再到小额实盘。"),
            ("AI 辅助写策略", "Jesse MCP 或只读 Freqtrade MCP + Cursor/Codex", "让 AI 帮你理解代码和跑回测，不直接控制资金。"),
            ("做市/网格/套利", "Hummingbot paper_trade + 小额 testnet", "更适合 order book、inventory、spread 和 executor。"),
            ("自己写底层执行", "CCXT + 自己的日志/风控/监控", "只适合程序员，必须自己补完整安全层。"),
            ("专业系统研究", "NautilusTrader", "工程更严谨，但学习门槛高。"),
        ],
        [1.25, 2.35, 2.9],
        size=8.0,
    )

    doc.add_heading("10. 实操安全清单", level=1)
    for item in [
        "只用小资金和独立子账户；API key 不开提现权限。",
        "任何框架先跑 backtest，再 dry-run/paper trading，再小额实盘。",
        "实盘前必须能解释：单笔最大亏损、每日最大亏损、最大持仓、强平价、手续费、滑点。",
        "所有订单、错误、余额变化、策略信号都要落日志。",
        "UI/API/MCP 不暴露公网；如果必须远程访问，用 VPN/SSH tunnel，不用弱密码。",
        "LLM/MCP 只用于辅助分析，不能无确认地下真实订单。",
        "不要下载闭源“套利 bot”；不要导入助记词；不要给陌生合约无限授权。",
    ]:
        number(doc, item)

    doc.add_heading("11. 来源链接", level=1)
    source(doc, "Freqtrade Backtesting", "https://docs.freqtrade.io/en/stable/backtesting/")
    source(doc, "Freqtrade REST API security warning", "https://www.freqtrade.io/en/stable/rest-api/")
    source(doc, "Freqtrade Strategy callbacks", "https://www.freqtrade.io/en/stable/strategy-callbacks/")
    source(doc, "Jesse docs and MCP", "https://docs.jesse.trade/")
    source(doc, "Jesse MCP introduction", "https://docs.jesse.trade/docs/mcp/")
    source(doc, "Hummingbot docs", "https://hummingbot.org/docs/")
    source(doc, "Hummingbot paper trade", "https://hummingbot.org/client/global-configs/paper-trade/")
    source(doc, "Hummingbot kill switch", "https://hummingbot.org/client/global-configs/kill-switch/")
    source(doc, "Hummingbot V2 controllers", "https://hummingbot.org/strategies/v2-strategies/controllers/")
    source(doc, "Hummingbot MCP", "https://hummingbot.org/mcp/")
    source(doc, "NautilusTrader docs", "https://nautilustrader.io/docs/")
    source(doc, "NautilusTrader execution", "https://nautilustrader.io/docs/latest/concepts/execution")
    source(doc, "NautilusTrader orders/risk checks", "https://nautilustrader.io/docs/latest/concepts/orders")
    source(doc, "CCXT manual", "https://github.com/ccxt/ccxt/wiki/manual")
    source(doc, "Freqtrade introspection MCP listing", "https://www.pulsemcp.com/servers/yalcin-freqtrade")
    source(doc, "OpenClaw local-first assistant", "https://openclawcn.com/en/")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
