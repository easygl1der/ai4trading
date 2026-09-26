from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


OUT = Path("/Users/yitwah/Documents/Codex/2026-06-07/video-youtube-video-open-document-recently/outputs/crypto_automation_tools_guide_2026.docx")


def rgb(value):
    return RGBColor.from_string(value)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_text(cell, text, bold=False, size=8.3, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def widths(table, values):
    for row in table.rows:
        for i, width in enumerate(values):
            set_width(row.cells[i], width)


def add_table(doc, headers, rows, col_widths, size=8.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        cell_text(table.rows[0].cells[i], header, bold=True, size=8.5, color="0B2545")
        shade(table.rows[0].cells[i], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, size=size)
    widths(table, col_widths)
    doc.add_paragraph("")
    return table


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def label_para(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(label + "：")
    r.bold = True
    p.add_run(text)


def add_callout(doc, title, text, fill="FFF4D6", color="5C3B00"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title + "：")
    r.bold = True
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.size = Pt(9)
    r.font.color.rgb = rgb(color)
    r2 = p.add_run(text)
    r2.font.name = "Arial"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r2.font.size = Pt(9)
    r2.font.color.rgb = rgb(color)
    doc.add_paragraph("")


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, bold in [
        ("Title", 24, "0B2545", False),
        ("Heading 1", 16, "0B2545", True),
        ("Heading 2", 13, "1F4D78", True),
        ("Heading 3", 11, "1F4D78", True),
    ]:
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(size)
        st.font.color.rgb = rgb(color)
        st.font.bold = bold
        st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph(style="Title")
    title.add_run("Crypto 自动化交易工具选择指南")
    sub = doc.add_paragraph()
    sub.add_run("Freqtrade、Hummingbot、CCXT、Jesse：自动化交易、回测、做市、交易所 API 怎么选").font.color.rgb = rgb("555555")
    doc.add_paragraph("整理日期：2026-06-07。资料来源优先采用官方文档；用途建议面向 crypto 小白和初级程序员。")

    add_callout(
        doc,
        "一句话结论",
        "如果你想先做低频策略回测和 dry-run，用 Freqtrade；如果你想做做市、跨所做市、流动性策略，用 Hummingbot；如果你想自己写交易所数据/下单接口，用 CCXT；如果你想做策略研究、回测可视化、AI/Notebook 辅助分析，用 Jesse。"
    )

    doc.add_heading("1. 工具选择矩阵", level=1)
    add_table(
        doc,
        ["需求", "首选", "备选", "不建议作为首选", "理由"],
        [
            ("自动化现货/合约策略", "Freqtrade", "Jesse", "CCXT 单独使用", "Freqtrade 有策略类、回测、dry-run、FreqUI/REST 控制；CCXT 只是 API 库，缺少完整风控和策略生命周期。"),
            ("回测", "Freqtrade / Jesse", "CCXT + 自写框架", "Hummingbot 新手直接回测", "Freqtrade 和 Jesse 都把回测作为核心能力；Jesse 可视化更强，Freqtrade 生态和配置更成熟。"),
            ("做市/挂单/套利", "Hummingbot", "CCXT 自写", "Freqtrade", "Hummingbot V2 controllers 明确面向 market making、多策略、长运行部署。"),
            ("交易所 API/行情抓取", "CCXT", "Freqtrade/Jesse 内置连接", "手写每家交易所 API", "CCXT 统一 public/private API，适合抓 OHLCV、订单簿、余额、下单。"),
            ("新手第一套实操", "Freqtrade dry-run", "Jesse backtest", "直接 Hummingbot 实盘做市", "新手先学习信号、回测、dry-run 和日志，不要一开始就暴露双边挂单库存风险。"),
            ("AI/Notebook 研究", "Jesse", "Freqtrade notebooks", "Hummingbot CLI", "Jesse docs 强调 research、optimization、Monte Carlo、MCP/AI-friendly workflow。"),
        ],
        [1.15, 1.05, 1.35, 1.45, 3.1],
        7.6,
    )

    doc.add_heading("2. 每个工具到底是什么", level=1)
    add_table(
        doc,
        ["工具", "定位", "最适合", "主要限制", "新手建议"],
        [
            ("Freqtrade", "Python crypto trading bot 框架", "低频/中频策略、技术指标信号、回测、dry-run、实盘自动交易", "不专门为做市设计；市场微结构、盘口策略、跨所库存管理需要额外工程", "先用 spot dry-run，学会策略、回测、日志、风险参数"),
            ("Hummingbot", "做市和 algo trading 框架", "market making、cross-exchange market making、grid、套利、DEX/CLMM 流动性", "配置和库存风险复杂；新手很容易被单边行情和手续费磨损", "只在模拟/小资金环境学习，先理解 spread、inventory、hedge"),
            ("CCXT", "多交易所 API 统一库", "行情、OHLCV、订单簿、余额、下单、撤单、私有账户接口", "不是策略框架；没有天然回测、仓位管理、风控、监控", "程序员必学，但不要用它直接裸奔实盘"),
            ("Jesse", "策略研究、回测、livetrade 框架", "多时间框架/多品种回测、可视化、导出、优化、Monte Carlo、AI/MCP 辅助", "部署栈包含 PostgreSQL/Redis；生态比 Freqtrade 小；做市不是主战场", "适合做研究和策略验证，尤其想看图和分析结果"),
        ],
        [0.9, 1.6, 2.3, 2.2, 1.55],
        7.8,
    )

    doc.add_heading("3. 推荐组合", level=1)
    label_para(doc, "组合 A：小白学习和第一套 bot", "Freqtrade + FreqUI。先写最简单 RSI/均线策略，下载历史数据，回测，然后 dry-run 至少两周。不要直接实盘。")
    label_para(doc, "组合 B：程序员做数据和执行层", "CCXT + 自己的日志/风控模块。适合做行情抓取、订单簿研究、交易所接口验证，但必须自己处理 rate limit、重试、订单状态、异常和风控。")
    label_para(doc, "组合 C：做市/套利/库存管理", "Hummingbot。先从 simple PMM 或 V2 scripts/controllers 学起，再看 XEMM、grid、LP rebalancer 等。重点不是“策略能跑”，而是库存、手续费和行情方向风险。")
    label_para(doc, "组合 D：策略研究和可视化", "Jesse。适合在 dashboard/Notebook 中比较多组回测，导出 CSV/JSON/Pine Script，做优化、Monte Carlo、rule significance testing。")

    doc.add_heading("4. Freqtrade：适合先学自动化交易和回测", level=1)
    bullet(doc, "官方策略文档说明：Freqtrade 策略是 Python class，使用 dataframe、OHLCV candles、indicators、entry/exit signals，并支持 long/short 的 entry/exit 表达。")
    bullet(doc, "回测需要先下载历史 OHLCV 数据；默认导出 backtest 结果到 user_data/backtest_results，并可导出 trades 做进一步分析。")
    bullet(doc, "官方明确区分 backtesting 和 dry-run：backtesting 快但容易被扭曲；dry-run 使用实时交易所数据但不真的下单，是更接近实盘的 forward testing。")
    bullet(doc, "Freqtrade 可以用 FreqUI、Telegram、REST API、Webhooks 控制或监控运行中的 bot。")
    bullet(doc, "Hyperopt 可以调参，但它本质上会反复跑 backtest；调得越多，越容易过拟合。")
    add_callout(doc, "Freqtrade 风险点", "不要买网上“高收益回测策略”直接跑。官方文档也提醒公开策略的漂亮回测常常不现实；必须做 lookahead/recursive analysis、dry-run 和小资金验证。", "FDEBD0", "6B3D00")

    doc.add_heading("5. Hummingbot：适合做市、套利和长运行策略", level=1)
    bullet(doc, "Hummingbot V2 把 scripts 作为学习/原型入口，把 controllers 作为生产级、可配置、可复用的长运行策略组件。")
    bullet(doc, "Controllers 适合多策略并行，例如在多个交易对上做 market making；它们通过 MarketDataProvider 读 OrderBook、Trades、Candles，并发出 ExecutorActions。")
    bullet(doc, "Hummingbot 的 V2 architecture 包含 PositionExecutor、DCAExecutor、GridExecutor、TWAPExecutor、XEMMExecutor、LPExecutor 等执行器。")
    bullet(doc, "启动 V2 策略前需要先 create config，然后用 start --v2 启动；Docker/headless 可以做自动启动。")
    add_callout(doc, "Hummingbot 风险点", "做市不是“挂买卖单稳赚 spread”。最大风险是 inventory 被单边行情打歪，或者手续费、滑点、撤单延迟、对冲失败吃掉 spread。新手只应先模拟或极小资金。", "FDEBD0", "6B3D00")

    doc.add_heading("6. CCXT：适合做交易所 API 的底层工具", level=1)
    bullet(doc, "CCXT 官方 manual 建议优先使用 unified methods，只有 unified method 不覆盖时再用 exchange-specific implicit methods。")
    bullet(doc, "Public API 可获取市场数据，如交易对、价格、订单簿、成交历史、ticker、OHLCV；通常不需要认证但受 rate limit 限制。")
    bullet(doc, "Private API 用于账户和交易：余额、下单、撤单、开仓/持仓、转账等，需要交易所 API key，部分交易所要求 KYC。")
    bullet(doc, "Unified API 的 params 常常仍然是交易所特定的，必须查对应交易所官方 API 文档。")
    add_callout(doc, "CCXT 风险点", "CCXT 只解决“怎么连接交易所”。它不帮你判断策略是否赚钱，也不自动处理完整风控。用 CCXT 实盘前，要自己实现幂等下单、撤单确认、失败重试、限速、日志、API key 最小权限。", "FDEBD0", "6B3D00")

    doc.add_heading("7. Jesse：适合策略研究、可视化和 AI 辅助", level=1)
    bullet(doc, "Jesse 官方首页强调可在没有 look-ahead bias 的情况下 backtest/livetrade 多时间框架和多品种。")
    bullet(doc, "Jesse 环境需要 Python 3.10-3.13、PostgreSQL、Redis；官方推荐 Docker 给初学者。")
    bullet(doc, "Backtest 结果可生成可视化图表，如累计收益 vs benchmark、drawdown、underwater plot、monthly returns、trade PnL distribution。")
    bullet(doc, "Exports 支持 CSV、JSON、TradingView Pine Script，适合做二次分析或把进出场点放到 TradingView 图上。")
    bullet(doc, "Research/Optimize 文档说明可用 Optuna/Ray 做参数搜索，并提示 training 好但 testing 崩溃是过拟合信号。")
    add_callout(doc, "Jesse 风险点", "Jesse 文档自己也写明不保证盈利、过去表现不代表未来、要注意 overfitting。它适合研究，不是赚钱按钮。", "FDEBD0", "6B3D00")

    doc.add_heading("8. 小白到实盘的建议路线", level=1)
    steps = [
        "第 1 周：只学交易所基础、订单簿、手续费、滑点、API key 权限。不要写自动下单。",
        "第 2 周：用 CCXT 只抓公开行情：ticker、OHLCV、order book。不要配置 private API。",
        "第 3-4 周：用 Freqtrade 写最简单策略，下载历史数据，跑 backtest，记录指标和失败场景。",
        "第 5-6 周：Freqtrade dry-run。让 bot 在实时行情中跑，但不真实下单；比较 dry-run 与 backtest 差异。",
        "第 7-8 周：用 Jesse 或 Freqtrade 做二次分析：drawdown、月度收益、交易分布、参数敏感性。",
        "第 9 周以后：如要做市，才开始 Hummingbot 模拟。先关注 inventory、spread、fees、hedge，不追求收益。",
        "实盘前：API key 只开交易权限，不开提现；小资金；单日最大亏损；日志；异常时能手动停机。",
    ]
    for step in steps:
        number(doc, step)

    doc.add_heading("9. 评分表", level=1)
    add_table(
        doc,
        ["维度", "Freqtrade", "Hummingbot", "CCXT", "Jesse"],
        [
            ("新手友好", "高", "中", "中-低", "中"),
            ("回测", "强", "中", "弱，需自写", "强"),
            ("实盘自动交易", "强", "强", "需自写", "中-强"),
            ("做市/套利", "弱-中", "强", "可自写但复杂", "弱-中"),
            ("交易所 API", "封装后可用", "封装后可用", "强", "封装后可用"),
            ("可视化/研究", "中", "中", "弱", "强"),
            ("部署复杂度", "中", "中-高", "低到高，看你写什么", "中-高"),
            ("最适合你现在", "第一优先", "第二阶段", "辅助学习 API", "研究/复盘辅助"),
        ],
        [1.45, 1.7, 1.7, 1.7, 1.7],
        8.0,
    )

    doc.add_heading("10. 我的推荐", level=1)
    add_callout(
        doc,
        "最适合你的路径",
        "先用 Freqtrade 做策略学习和 dry-run；同时用 CCXT 学市场数据和 API 基础；等你理解 spread、库存和手续费后，再用 Hummingbot 做 market making；如果你想更系统地做回测图表、参数优化和 AI 辅助研究，再引入 Jesse。",
        "E8F4FD",
        "0B2545",
    )
    bullet(doc, "不要一开始就做 Hummingbot 实盘做市。做市看起来机械，但风险比普通低频现货策略更隐蔽。")
    bullet(doc, "不要直接用 CCXT 自写全自动实盘 bot，除非你已经有日志、错误恢复、风控、订单状态机和 kill switch。")
    bullet(doc, "不要相信任何“回测 300% 收益 + 低回撤”的公开策略；先检查 lookahead、手续费、滑点、成交量、样本外表现。")

    doc.add_heading("11. 官方资料链接", level=1)
    links = [
        "Freqtrade Strategy Quickstart: https://docs.freqtrade.io/en/stable/strategy-101/",
        "Freqtrade Backtesting: https://docs.freqtrade.io/en/stable/backtesting/",
        "Freqtrade Hyperopt: https://docs.freqtrade.io/en/stable/hyperopt/",
        "Hummingbot Documentation: https://hummingbot.org/docs/",
        "Hummingbot V2 Strategy Architecture: https://hummingbot.org/strategies/v2-strategies/",
        "Hummingbot Controllers: https://hummingbot.org/strategies/v2-strategies/controllers/",
        "Hummingbot Start/Stop Strategy: https://hummingbot.org/client/start-stop/",
        "CCXT Manual: https://github.com/ccxt/ccxt/wiki/manual",
        "Jesse Docs: https://docs.jesse.trade/",
        "Jesse Getting Started: https://docs.jesse.trade/docs/getting-started/",
        "Jesse Backtest Exports: https://docs.jesse.trade/docs/backtest/exports",
        "Jesse Result Charts: https://docs.jesse.trade/docs/backtest/charts",
        "Jesse Optimize Research Function: https://docs.jesse.trade/docs/research/optimize",
    ]
    for link in links:
        bullet(doc, link)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Crypto automation tools guide | generated 2026-06-07")
    run.font.size = Pt(8)
    run.font.color.rgb = rgb("666666")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
