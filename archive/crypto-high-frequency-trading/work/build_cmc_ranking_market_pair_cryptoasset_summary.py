from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path("/Users/yitwah/Documents/Codex/2026-06-07/video-youtube-video-open-document-recently")
OUT = BASE / "outputs/cmc_ranking_market_pair_cryptoasset_summary_zh.docx"


SOURCES = [
    (
        "Ranking (Market Pair, Cryptoasset)",
        "https://support.coinmarketcap.com/hc/en-us/articles/360043836851-Ranking-Market-Pair-Cryptoasset",
    ),
    (
        "Volume & Open Interest (Market Pair, Cryptoasset, Exchange, Aggregate)",
        "https://support.coinmarketcap.com/hc/en-us/articles/360043395912-Volume-Open-Interest-Market-Pair-Cryptoasset-Exchange-Aggregate",
    ),
    (
        "Liquidity Score (Market Pair, Exchange)",
        "https://support.coinmarketcap.com/hc/en-us/articles/360043836931-Liquidity-Score-Market-Pair-Exchange",
    ),
    (
        "Confidence Indicator (Market Pair)",
        "https://support.coinmarketcap.com/hc/en-us/articles/360044481772-Confidence-Indicator-Market-Pair",
    ),
    (
        "Market Data & Cryptoasset Rank",
        "https://support.coinmarketcap.com/hc/en-us/articles/360034116491-Market-Data-Cryptoasset-Rank",
    ),
    (
        "Metric Methodologies",
        "https://support.coinmarketcap.com/hc/en-us/sections/360008888252-Metric-Methodologies",
    ),
]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_font(run, size=10, color="111827", bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches: float):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(int(w * 1440) for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for i, width in enumerate(widths):
            set_cell_width(row.cells[i], width)
            set_cell_margins(row.cells[i])
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def cell_text(cell, text, *, bold=False, size=8.4, color="111827", fill=None, align=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align:
        p.alignment = align
    r = p.add_run(str(text))
    set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, *, size=8.1, header_fill="E8EEF5"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell_text(
            table.rows[0].cells[i],
            header,
            bold=True,
            size=8.4,
            color="0B2545",
            fill=header_fill,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(row) > 2 else None
            cell_text(cells[i], value, size=size, align=align)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, title, body, *, fill="E8F4FD", color="0B2545"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(f"{title}: ")
    set_font(r, size=9.4, color=color, bold=True)
    r2 = p.add_run(body)
    set_font(r2, size=9.4, color=color)
    set_table_geometry(table, [6.34])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def para(doc, text, *, size=10, color="111827", after=6, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.14
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, size=size, color=color, bold=True)
        text = text[len(bold_prefix) :]
    r = p.add_run(text)
    set_font(r, size=size, color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=9.6)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=9.6)


def source_line(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(f"{label}: {url}")
    set_font(r, size=8.1, color="4B5563")


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("CoinMarketCap 官方方法论中文总结 | 仅供学习，不构成投资建议")
    set_font(r, size=8, color="6B7280")


def setup_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.76)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.42)
    sec.footer_distance = Inches(0.38)
    add_footer(sec)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb("111827")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.14

    for style_name, size, color, bold, before, after in [
        ("Title", 21.5, "0B2545", False, 0, 6),
        ("Subtitle", 9.4, "4B5563", False, 0, 8),
        ("Heading 1", 15, "0B2545", True, 13, 6),
        ("Heading 2", 12.3, "1F4D78", True, 9, 4),
        ("Heading 3", 10.8, "1F4D78", True, 7, 3),
        ("List Bullet", 9.6, "111827", False, 0, 3),
        ("List Number", 9.6, "111827", False, 0, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.12


def add_title(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("CoinMarketCap 排名、交易对与 Cryptoasset：中文新手总结")
    set_font(r, size=21.5, color="0B2545")

    sub = doc.add_paragraph(style="Subtitle")
    r = sub.add_run(
        "资料对象：CoinMarketCap 官方支持/方法论页面《Ranking (Market Pair, Cryptoasset)》"
        "｜用户给定更新时间：2026-01-21｜整理日期：2026-06-07"
    )
    set_font(r, size=9.4, color="4B5563")


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    add_title(doc)

    add_callout(
        doc,
        "一句话结论",
        "CMC 的排名不是把交易所自己报出来的 24 小时成交量从大到小排序。对 market pair，CMC 会同时看 Reported Volume、Liquidity Score 和 Web Traffic Factor，并用 Confidence Indicator 提醒交易量是否可信；对 cryptoasset，CMC 还会看市值可验证性、流动性、价差、资料完整度和反操纵因素。",
    )
    add_callout(
        doc,
        "新手最容易踩的坑",
        "看到某个币或交易对“成交量很大”不等于它真的容易买卖，也不等于项目更安全。成交量可能来自无手续费、返佣、交易挖矿或刷量；真正下单时，你关心的是能否以接近预期的价格成交，以及这些数据是否被 CMC 认为足够可信。",
        fill="FFF4D6",
        color="5C3B00",
    )

    doc.add_heading("1. 三个关键词先讲清楚", level=1)
    add_table(
        doc,
        ["词", "新手版解释", "为什么影响排名/交易判断"],
        [
            (
                "Cryptoasset",
                "一个被 CMC 跟踪的加密资产，比如 coin、token、稳定币、包装资产或质押衍生资产。",
                "资产本身的排名要看市值与供应量是否可验证，还要避免 wrapped/staked 等结构造成市值重复计算。",
            ),
            (
                "Market pair",
                "一个具体交易市场，例如 BTC/USDT、ETH/USD。它包含 base asset 和 quote asset，并且通常发生在某个交易所。",
                "同一个币在不同交易所、不同报价币种下可能流动性完全不同，所以交易对排名不能只看币的名气。",
            ),
            (
                "Ranking",
                "CMC 给交易对、资产或交易所排序的规则集合。",
                "排序用于帮用户筛选，但 CMC 明确不会公开全部内部阈值，以降低项目或交易所“按答案作弊”的空间。",
            ),
        ],
        [1.2, 2.6, 2.54],
        size=7.9,
    )
    para(
        doc,
        "从阅读顺序看，先理解 market pair，再理解 cryptoasset rank：交易对回答“去哪里交易更顺”，资产排名回答“这个资产在 CMC 体系里是否有足够可验证、市值与流动性基础”。",
    )

    doc.add_heading("2. Market Pair Rank：为什么不能只看 reported volume", level=1)
    para(
        doc,
        "CMC 官方文章说，交易量膨胀是它自 2018 年以来持续处理的问题：部分交易所可能上报被夸大的交易量，用来制造合法性和流动性错觉。CMC 因此把 market pair ranking 从单纯依赖成交量，改成同时考虑 Reported Volume、Liquidity Score 和 Web Traffic Factor，并引入 Confidence Indicator。",
    )
    add_table(
        doc,
        ["指标", "它大概衡量什么", "新手应该怎么用"],
        [
            (
                "Reported Volume",
                "交易所直接上报的 24 小时成交量，再按 CMC 参考价格换算成美元口径。",
                "只能当原始输入，不要直接当作“真实热度”或“能顺利成交”的证明。",
            ),
            (
                "Liquidity Score",
                "CMC 根据订单簿深度和不同金额立即买卖会产生的滑点，给出 0-1000 的流动性分数。",
                "更接近“我真的下单会不会被滑点伤到”。分数高通常意味着买卖更顺，但它仍是量化指标，不是交易所信用背书。",
            ),
            (
                "Web Traffic Factor",
                "用交易所网页流量近似反映交易者数量或市场参与度。",
                "帮助识别“成交量很大但看起来没有足够用户支撑”的异常情况。",
            ),
            (
                "Confidence Indicator",
                "CMC 用模型估计每个交易对合理交易量，并标记 reported volume 的可信程度。",
                "优先关注 High/Moderate；Low 代表 CMC 对该交易对上报量的信心较低，应额外谨慎。",
            ),
        ],
        [1.25, 2.65, 2.44],
        size=7.9,
    )
    add_callout(
        doc,
        "Reported volume 的核心问题",
        "成交量是“发生了多少交易”的口径，但流动性是“你下单能不能以接近期望价格成交”的口径。一个市场可以上报很大的成交量，却在订单簿两边没有足够真实深度，导致你买入或卖出时出现明显滑点。",
        fill="F4F6F9",
        color="1F3A5F",
    )

    doc.add_heading("3. Reported、Adjusted 与“Verified”交易量怎么区分", level=1)
    add_table(
        doc,
        ["口径", "CMC 官方语境", "怎么理解"],
        [
            (
                "Reported Volume",
                "交易所全部现货市场上报的 24 小时交易量。",
                "最宽口径，容易被交易激励、无手续费或刷量行为污染。",
            ),
            (
                "Adjusted Volume",
                "CMC 官方定义为排除无手续费市场和 transaction mining 市场后的现货交易量。",
                "更保守，但不是保证“完全真实”；它只是剔除一类高风险交易量来源。",
            ),
            (
                "Cryptoasset Volume",
                "某资产在所有交易所的 24 小时现货交易量汇总；CMC 会排除部分易受 wash trading 影响的交易对。",
                "看资产热度时更适合，但仍要结合交易对分布和流动性。",
            ),
            (
                "Verified Volume",
                "目标文章没有把它定义成一个独立官方指标；官方更常说的是 CMC-verified market cap/supply，以及 Confidence 对 reported volume 的可信判断。",
                "阅读时不要把“verified volume”当作绝对真实成交量。更稳妥的说法是：看 adjusted volume、confidence、liquidity score 和可验证供应信息的组合。",
            ),
        ],
        [1.25, 2.8, 2.29],
        size=7.75,
    )

    doc.add_heading("4. Liquidity Score：它比成交量更贴近真实交易体验", level=1)
    para(
        doc,
        "CMC 的 Liquidity Score 是 0 到 1000 的分数。官方解释中，1000 代表最有流动性的市场，0 代表最不流动。它重点模拟不同订单金额立即买入和卖出时的滑点，并优先关注多数零售用户更相关的订单规模。",
    )
    for item in [
        "滑点越低，说明你实际成交价格越接近下单前看到的价格。",
        "买盘和卖盘两边都要有足够订单；只有一边深度好，并不代表你进出都顺。",
        "高分市场通常更适合交易，但 CMC 也提醒：流动性不是选择交易所或交易对的唯一标准。",
        "Liquidity Score 与 volume 是两个概念：前者衡量市场可成交质量，后者统计资产换手量。",
    ]:
        bullet(doc, item)
    add_table(
        doc,
        ["场景", "reported volume 看起来", "Liquidity/Confidence 可能揭示"],
        [
            ("正常活跃市场", "成交量高", "订单簿厚、滑点低、信心高，比较适合进一步研究。"),
            ("刷量或激励市场", "成交量很高", "交易者数量、订单簿深度或信心不足，可能只是数字好看。"),
            ("小币冷门交易对", "成交量低或不稳定", "下单金额稍大就滑点明显，卖出时尤其危险。"),
            ("价差异常市场", "短期成交量突出", "跨交易所价格差明显，可能反映数据异常或真实套利成本很高。"),
        ],
        [1.25, 1.55, 3.54],
        size=7.9,
    )

    doc.add_heading("5. Cryptoasset Rank：资产进入 Top 200 不是只靠市值", level=1)
    para(
        doc,
        "目标文章说明，Top 200 Cryptoasset Rank 的资格由 market capitalization 以及多项因素共同决定。CMC 特别强调供应信息可验证、无市值重复计算、充分流动性/交易活动、正常 bid-ask spread、资料评分、没有重大跨交易所价格偏差，以及在至少三个具备较好数据/API/合规或社区属性的交易所交易。",
    )
    add_table(
        doc,
        ["排名状态", "最低条件/含义", "新手解读"],
        [
            (
                "Top 200 Rank",
                "至少有 CMC-verified market cap，并满足目标文章列出的核心要求。",
                "更强调数据可验证性与市场质量，不只是“市值算出来很大”。",
            ),
            (
                "Top 201 and beyond",
                "至少有 CMC-verified market cap，但不需要满足 Top 200 的全部要求。",
                "可能市值已可验证，但流动性、资料或市场结构不足以进入前 200 资格体系。",
            ),
            (
                "Unranked",
                "没有 CMC-verified market cap 的 cryptoasset，按 24 小时交易量排序。",
                "不是“没有价值”的同义词，但说明 CMC 没有把它放进正常市值排名体系。",
            ),
            (
                "Rehypothecated crypto",
                "单独排名，以减少榜单干扰并避免市值重复计算。",
                "wrapped、staked、衍生包装类资产要格外注意重复计算和赎回结构。",
            ),
        ],
        [1.35, 2.58, 2.41],
        size=7.75,
    )

    doc.add_heading("6. 虚假交易量 / Wash Trading 风险", level=1)
    para(
        doc,
        "CMC 的 Volume 页面提到，无交易费或提供显著交易激励的市场更容易出现 wash trading，导致 reported volume 被人为放大。对新手来说，wash trading 可以理解为“看起来成交很多，但很多成交未必代表真实买卖需求”。",
    )
    for item in [
        "可能目的：制造项目热度、吸引上币或做市合作、误导散户以为市场很深。",
        "常见危险信号：成交量异常高但网页流量、盘口深度、价差和交易所可信度跟不上。",
        "实际风险：你按成交量以为容易退出，真正卖出时却找不到足够买盘，价格被自己砸下去。",
        "防御方式：同时看 adjusted volume、liquidity score、confidence indicator、交易所质量、价差和订单簿，而不是只看一个 24h volume 数字。",
    ]:
        bullet(doc, item)

    doc.add_heading("7. 新手查看 CMC 页面时的 6 步检查", level=1)
    for item in [
        "先确认自己看的是 cryptoasset 页面还是具体 market pair 页面，不要把资产热度和某个交易市场的流动性混为一谈。",
        "看交易对排名时，优先找 reported volume、liquidity score、web traffic/交易所质量和 confidence 都相对协调的市场。",
        "遇到高 reported volume 但低 confidence 的交易对，默认把它当作高风险信号，而不是便宜机会。",
        "看资产排名时，检查它是否有 CMC-verified market cap，以及是否因供应、包装资产、质押资产或价格异常被限制排名。",
        "准备交易前，先估算自己的订单金额与该交易对深度是否匹配；小额顺利不代表大额也顺利。",
        "把 CMC 排名当筛选工具，不当投资结论；最终仍要看项目基本面、合约/链上风险、交易所风险和自身风控。",
    ]:
        number(doc, item)

    doc.add_heading("8. 官方来源与取材边界", level=1)
    para(
        doc,
        "本文优先依据 CoinMarketCap 官方支持/方法论页面。目标文章可访问，页面标题为 Ranking (Market Pair, Cryptoasset)，用户给定更新时间为 2026-01-21；辅助页面用于补充 volume、liquidity score 和 confidence indicator 的定义。由于 CMC 不公开 ranking 算法的内部阈值和完整权重，本文只做新手解释，不推断未公开参数。",
        size=9.6,
    )
    for label, url in SOURCES:
        source_line(doc, label, url)

    doc.core_properties.title = "CoinMarketCap 排名、交易对与 Cryptoasset：中文新手总结"
    doc.core_properties.subject = "CoinMarketCap Ranking / Market Pair / Cryptoasset 官方方法论中文总结"
    doc.core_properties.author = "OpenAI Codex"
    doc.save(OUT)


if __name__ == "__main__":
    build()
