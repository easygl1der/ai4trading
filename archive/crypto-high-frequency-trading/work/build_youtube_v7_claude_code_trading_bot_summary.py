from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path("/Users/yitwah/Documents/Codex/2026-06-07/video-youtube-video-open-document-recently")
OUT = BASE / "outputs/youtube_v7_claude_code_trading_bot_notes_zh.docx"


VIDEO = {
    "title": "How To Actually Build a Trading Bot With Claude Code (Fully Automated)",
    "channel": "AI Pathways",
    "url": "https://www.youtube.com/watch?v=y_bsjZThP0o",
    "published": "2026-04-10",
    "duration": "34:18",
    "video_id": "y_bsjZThP0o",
}


CHAPTERS = [
    ("00:00", "Claude Code Trading Bot", "展示最终 dashboard：市场 regime、置信度、资产规模、买力、交易记录、风控状态。"),
    ("03:28", "Trading Bot Structure", "把系统拆成 brain、allocation、safety、brokerage、dashboard 五层。"),
    ("10:09", "Project Scaffolding", "用 Claude Code 建项目骨架、目录、测试入口和文档。"),
    ("13:15", "HMM Engine", "用 Hidden Markov Model 识别市场 regime，如 crash、bear、neutral、bull、euphoria。"),
    ("16:28", "Allocation Strategies", "按 regime 调整仓位和策略参数，市场越动荡，风险暴露越低。"),
    ("18:34", "Walk Forward Backtesting", "用 walk-forward validation 避免只在历史区间拟合得好。"),
    ("21:39", "Risk Management", "加入 circuit breakers、回撤限制、仓位限制和系统停机条件。"),
    ("24:59", "Connect Broker", "连接 Alpaca，优先使用 paper account；API key 放入本地 .env，不交给 Claude。"),
    ("29:33", "Main Loop & Orchestrator", "把配置、账户校验、训练、风控、数据流、下单和错误处理串起来。"),
]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_font(run, size=10, color="000000", bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            set_cell_width(row.cells[i], width)


def cell_text(cell, text, bold=False, size=8.2, color="000000", fill=None, align=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.10
    if align:
        p.alignment = align
    r = p.add_run(str(text))
    set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, size=7.9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        cell_text(hdr[i], header, bold=True, size=8.4, color="0B2545", fill="E8EEF5", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if widths[i] <= 1.05 else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cells[i], value, size=size, align=align)
    set_table_widths(table, widths)
    doc.add_paragraph("")
    return table


def add_callout(doc, title, body, fill="FFF4D6", color="5C3B00"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 7.05)
    shade(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=150, end=150)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.13
    r = p.add_run(f"{title}：")
    set_font(r, size=9.3, color=color, bold=True)
    r2 = p.add_run(body)
    set_font(r2, size=9.3, color=color)
    doc.add_paragraph("")


def para(doc, text, size=10, color="000000"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.13
    r = p.add_run(text)
    set_font(r, size=size, color=color)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.24 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=9.6)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=9.6)


def source(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{label}: {value}")
    set_font(r, size=8.2, color="555555")


def setup_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    sec.header_distance = Inches(0.38)
    sec.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13

    for name, size, color, bold, before, after in [
        ("Title", 22.5, "0B2545", False, 0, 8),
        ("Subtitle", 9.6, "555555", False, 0, 8),
        ("Heading 1", 15.2, "0B2545", True, 13, 6),
        ("Heading 2", 12.4, "1F4D78", True, 9, 5),
        ("Heading 3", 10.8, "1F4D78", True, 7, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(9.6)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.12


def add_footer(doc):
    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.paragraph_format.space_before = Pt(0)
        footer.paragraph_format.space_after = Pt(0)
        r = footer.add_run("AI Pathways Claude Code Trading Bot 中文记录")
        set_font(r, size=7.8, color="777777")


def build():
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Claude Code 自动化交易 Bot 视频中文记录")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(
        f"视频：{VIDEO['title']}｜频道：{VIDEO['channel']}｜时长：{VIDEO['duration']}｜发布：{VIDEO['published']}｜整理日期：2026-06-07"
    )

    add_callout(
        doc,
        "资料边界",
        "已通过 BrowserOS YouTube 工具读取视频详情、搜索确认和字幕。本文是面向程序员与 crypto 新手的结构化中文学习记录，不是逐字稿翻译；涉及 HMM、回测、券商 API、Alpaca、Claude Code 的内容均按视频字幕和描述章节归纳。",
        "E8F4FD",
        "0B2545",
    )
    add_callout(
        doc,
        "先看风险",
        "AI 生成交易 bot 不是赚钱保证。AI 可能写错代码、忽略异常、误用 API、泄露密钥、让策略过拟合历史数据，或者把 paper trading 中看似可行的逻辑错误搬到实盘。任何真实资金接入前，都必须人工审查、单元测试、回测、模拟盘、限额、监控和可立即停机的风控。",
    )

    doc.add_heading("1. 视频主旨", level=1)
    para(
        doc,
        "这支视频演示如何用 Claude Code 从零构建一个“看起来完整”的自动化交易系统：先识别市场状态，再根据状态调整仓位和策略，最后连接券商 API 执行订单，并用 dashboard 观察信号、仓位、风控和历史交易。",
    )
    para(
        doc,
        "它的核心卖点不是某个具体指标，而是把交易 bot 当成工程系统：模型、仓位、风控、回测、券商接入、主循环、日志和监控要分层开发，并且每一层都要独立测试。视频里强调先 paper trade 至少一个月，再考虑真实账户。",
    )

    doc.add_heading("2. 时间线与章节摘要", level=1)
    add_table(doc, ["时间", "章节", "中文记录重点"], CHAPTERS, [0.8, 2.1, 4.35], size=7.8)

    doc.add_heading("3. Bot 模块划分", level=1)
    add_table(
        doc,
        ["模块", "视频中的作用", "新手理解", "必须验证的点"],
        [
            ("Brain / HMM", "用 Hidden Markov Model 识别 crash、bear、neutral、bull、euphoria 等 market regime。", "先判断市场状态，再决定策略，而不是永远跑同一套规则。", "训练数据、特征、状态解释、置信度、极端行情下是否失真。"),
            ("Allocation", "按 regime 调整仓位、交易方向、策略参数和风险暴露。", "牛市、熊市、震荡市使用不同仓位，不把所有环境当成一样。", "最大仓位、单笔风险、再平衡频率、手续费与滑点。"),
            ("Safety", "独立于模型的 circuit breakers、回撤限制、杠杆状态和停机逻辑。", "风控要硬编码，不依赖 AI 或模型自己“聪明”。", "日内最大亏损、连续错误、账户异常、API 失败后是否停机。"),
            ("Brokerage", "连接 Alpaca 账户，读取账户状态并提交测试订单。", "交易 API 是资金入口，paper 和 live 必须严格分开。", ".env 密钥、本地权限、paper endpoint、订单状态确认。"),
            ("Dashboard", "用 Streamlit 展示 regime、资产、交易记录、风控状态和图表。", "UI 不是装饰，是实盘前监控和复盘入口。", "日志完整性、实时性、异常告警、能否追溯每次下单原因。"),
        ],
        [1.05, 2.05, 2.05, 2.2],
        size=7.3,
    )

    doc.add_heading("4. Claude Code / AI 辅助构建流程", level=1)
    for item in [
        "先给 Claude Code 一个 master document 或分阶段 prompt，让它理解系统目标、目录结构、模块边界和测试要求。",
        "从 project scaffolding 开始，建立目录、配置、依赖、测试文件和 README，而不是直接让 AI 写一整个交易脚本。",
        "逐层实现：HMM engine、allocation、walk-forward backtesting、risk manager、broker integration、orchestrator、dashboard。",
        "每一层写完都运行测试；视频后段提到系统通过了 134 个测试，再继续接 dashboard。",
        "敏感信息不交给 Claude：API key、secret key 由人手动写入本地 .env，并确保 .gitignore 忽略 .env。",
        "AI 可以生成代码和文档，但人必须检查交易逻辑、异常处理、API 权限和是否真的符合策略意图。",
    ]:
        bullet(doc, item)

    doc.add_heading("5. 自动化开发边界", level=1)
    add_callout(
        doc,
        "边界原则",
        "Claude Code 适合做工程助手，不适合直接成为资金负责人。它可以 scaffold、写测试、补文档、重构、接 API，但交易规则、资金限额、实盘开关和密钥管理应由人控制。",
        "F4F6F9",
        "1F3A5F",
    )
    add_table(
        doc,
        ["AI 可以做", "人必须负责", "原因"],
        [
            ("生成项目骨架、类、函数、测试、README。", "确认目录结构、命名、依赖和运行环境。", "AI 常会假设不存在的包、参数或 API 行为。"),
            ("根据 prompt 实现 HMM、回测、风险管理器和 broker adapter。", "审查数学假设、数据泄漏、边界条件和失败路径。", "交易系统错一次可能就是资金损失，不是普通 UI bug。"),
            ("写自动化测试和日志。", "选择真正有意义的测试样本和验收标准。", "测试覆盖了代码路径，不等于覆盖了市场风险。"),
            ("解释报错、提出修复、补 dashboard。", "决定是否上线、是否加资金、是否打开实盘。", "上线是风险决策，不是代码生成任务。"),
        ],
        [2.25, 2.45, 2.25],
        size=7.7,
    )

    doc.add_heading("6. 交易 API 与账户安全", level=1)
    para(doc, "视频使用 Alpaca 作为券商接入示例，并建议先使用 paper account。流程是从 Alpaca 账户复制 base URL、API key、secret key，把 key 写入本地 .env，再让 broker integration 从环境变量读取。")
    for item in [
        "不要把 API key 或 secret key 贴进 Claude Code 聊天、截图、GitHub、日志或 issue。",
        "先用 paper trading endpoint；paper=true 与 live endpoint 要在配置层明确分开。",
        "API key 只给最小必要权限；新手尤其不要开提现或资金转出权限。",
        "所有下单接口都要处理：提交成功但未成交、部分成交、拒单、超时、重复请求、市场关闭、价格精度错误。",
        "网络/API 异常后，默认策略应是停止新增风险，而不是继续尝试加仓。",
    ]:
        bullet(doc, item)

    doc.add_heading("7. 回测、模拟、实盘的区别", level=1)
    add_table(
        doc,
        ["阶段", "目的", "能证明什么", "不能证明什么"],
        [
            ("历史回测", "用历史数据检查策略逻辑和风险指标。", "能暴露明显亏损、未来函数、参数脆弱性。", "不能证明未来赚钱，尤其不能覆盖真实滑点和订单失败。"),
            ("Walk-forward", "按时间滚动训练/验证，降低过拟合。", "能比单次回测更接近“未来未知”的测试方式。", "仍然依赖历史样本，无法穷尽 regime 切换。"),
            ("Paper trading", "用真实行情和模拟资金跑完整链路。", "能检查 API、订单、日志、dashboard 和实时流程。", "不能完全模拟真实成交、心理压力和资金规模影响。"),
            ("小额实盘", "用极低风险资金验证真实执行。", "能看到真实手续费、滑点、订单状态和账户行为。", "不能保证扩大资金后仍有效。"),
            ("规模化实盘", "在严格风控下持续运行。", "只能逐步积累证据。", "没有任何阶段能给出稳赚保证。"),
        ],
        [1.05, 2.0, 2.25, 2.25],
        size=7.3,
    )

    doc.add_heading("8. 代码生成风险清单", level=1)
    for item in [
        "未来函数：回测代码误用未来价格、未来成交量或未来 regime。",
        "过拟合：参数只对一个 ticker、一个年份、一个牛市窗口有效。",
        "单位错误：价格、数量、百分比、杠杆、base/quote currency 混淆。",
        "订单重复：超时后重试，没有查询已有订单，导致重复买入或卖出。",
        "异常路径空白：API rate limit、市场关闭、余额不足、交易对不可用、数据缺失没有处理。",
        "密钥泄露：.env 没有进 .gitignore，或把 key 写进代码、日志、prompt。",
        "测试虚假安全感：通过很多测试，但测试只验证函数返回，不验证交易假设。",
    ]:
        bullet(doc, item)

    doc.add_heading("9. 风控与安全底线", level=1)
    add_callout(
        doc,
        "实盘前最低门槛",
        "至少完成：多市场/多年份回测、walk-forward、paper trading 一个月以上、人工复盘每笔交易、API 权限最小化、单笔/单日/总仓位限额、异常停机、日志可追溯、密钥轮换和撤销流程。",
        "FCE8E6",
        "9B1C1C",
    )
    add_table(
        doc,
        ["风控项", "建议实现", "为什么重要"],
        [
            ("单笔风险", "限制每笔订单占账户净值比例，并按 stop loss 计算实际风险。", "避免一次错误订单毁掉账户。"),
            ("日内亏损", "达到日内最大回撤后停止交易。", "防止系统在异常行情中连续亏损。"),
            ("总仓位", "按 asset、sector、regime 限制最大暴露。", "避免多个信号实际押同一方向。"),
            ("熔断", "连续报错、API 超时、数据断流、账户余额异常时关闭新增订单。", "交易系统必须知道什么时候不交易。"),
            ("人工确认", "从 paper 到 live、从小额到加仓，都要人工批准。", "资金风险不能交给自动代码自行升级。"),
            ("审计日志", "记录数据、信号、模型状态、订单请求、订单回执、异常。", "没有日志无法复盘，也无法确认责任边界。"),
        ],
        [1.1, 3.0, 2.85],
        size=7.6,
    )

    doc.add_heading("10. 给程序员与 crypto 新手的落地路线", level=1)
    for item in [
        "第一周：只读视频和文档，复现目录结构，不接真实账户。",
        "第二周：实现数据层、回测层和一个极简单策略，重点检查手续费、滑点、时间戳和测试。",
        "第三周：加入 regime 或风险状态，但先把它当观察变量，不让它自动下大单。",
        "第四周：接 paper trading，运行完整主循环和 dashboard，每天复盘所有信号。",
        "一个月后：若 paper 逻辑稳定，也只用极小资金做实盘验证，并保留随时停机的人工开关。",
    ]:
        number(doc, item)

    doc.add_heading("11. 来源", level=1)
    source(doc, "YouTube 视频", VIDEO["url"])
    source(doc, "频道", VIDEO["channel"])
    source(doc, "视频 ID", VIDEO["video_id"])
    source(doc, "BrowserOS YouTube transcript", "读取成功；本文使用字幕内容做中文归纳，不转载完整逐字稿。")
    source(doc, "视频描述链接", "AI Trading Community: https://www.skool.com/aipathways；1:1 AI Trading Program: https://www.aipathways.io/investing")

    add_footer(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
