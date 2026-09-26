from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path("/Users/yitwah/Documents/Codex/2026-06-07/video-youtube-video-open-document-recently")
OUT = BASE / "outputs/cmc_how_can_i_buy_coins_tokens_summary_zh.docx"

SOURCES = [
    (
        "CoinMarketCap Help Center: How can I buy coins/tokens?",
        "https://support.coinmarketcap.com/hc/en-us/articles/360013803852-How-can-I-buy-coins-tokens",
    ),
    (
        "CoinMarketCap Academy: How to Buy and Sell Crypto on CoinMarketCap",
        "https://coinmarketcap.com/academy/article/how-to-buy-and-sell-crypto-on-coinmarketcap",
    ),
    ("CoinMarketCap Disclaimer", "https://coinmarketcap.com/disclaimer/"),
]


def rgb(value):
    return RGBColor.from_string(value)


def set_font(run, size=10, color="202124", bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)

    for row in table.rows:
        for i, width in enumerate(widths):
            set_cell_width(row.cells[i], width)
            set_cell_margins(row.cells[i])


def cell_text(cell, text, bold=False, size=8.4, color="202124", fill=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    if fill:
        shade(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(str(text))
    set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell_text(table.rows[0].cells[i], header, bold=True, size=8.6, color="0B2545", fill="E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, size=size)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def add_callout(doc, title, body, fill="E8F4FD", color="0B2545"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=170, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{title}：")
    set_font(r, size=9.4, color=color, bold=True)
    r2 = p.add_run(body)
    set_font(r2, size=9.4, color=color)
    set_table_geometry(table, [6.5])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def para(doc, text, size=10, color="202124", after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.14
    r = p.add_run(text)
    set_font(r, size=size, color=color)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=9.5)
    return p


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=9.5)
    return p


def source_line(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(f"{label}: {url}")
    set_font(r, size=8.2, color="5F6368")


def set_borders_off(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def add_key_facts(doc):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    facts = [
        ("官方文章", "How can I buy coins/tokens?"),
        ("官方更新时间", "2026-05-07 19:09（用户给定并与检索结果一致）"),
        ("核心结论", "CMC 是市场数据与交易入口界面，不是交易所，也不是钱包。"),
        ("适用读者", "第一次用 CMC 找买币入口、还不熟悉钱包/DEX/授权/滑点的新手。"),
    ]
    for row, (k, v) in zip(table.rows, facts):
        cell_text(row.cells[0], k, bold=True, size=8.5, color="0B2545", fill="F2F4F7")
        cell_text(row.cells[1], v, size=8.5)
    set_table_geometry(table, [1.55, 4.95])
    doc.add_paragraph("")


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb("202124")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.14

    for name, size, color, bold, before, after in [
        ("Title", 22.5, "0B2545", False, 0, 8),
        ("Heading 1", 15.5, "0B2545", True, 13, 6),
        ("Heading 2", 12.5, "1F4D78", True, 9, 5),
        ("Heading 3", 11, "1F4D78", True, 7, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(9.5)
        style.paragraph_format.space_after = Pt(3.5)
        style.paragraph_format.line_spacing = 1.12


def add_footer(doc):
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("CoinMarketCap 帮助文章中文总结｜仅作学习参考，不构成投资建议")
    set_font(r, size=8.2, color="7A7A7A")


def build():
    doc = Document()
    setup_styles(doc)
    add_footer(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("CoinMarketCap 买币帮助文章中文总结")

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(8)
    r = sub.add_run("面向 crypto 小白｜整理日期：2026-06-07｜官方文章更新时间：2026-05-07")
    set_font(r, size=9.4, color="5F6368")

    add_key_facts(doc)

    add_callout(
        doc,
        "先记住一句话",
        "CoinMarketCap 可以帮你查看价格、进入币种详情页、找到市场或通过 DEX Mode 发起链上 swap；但 CMC 不替你保管资产、不替你买卖、不推荐你买哪一个币。真正签名和承担后果的是你的钱包和你自己。",
    )
    add_callout(
        doc,
        "资料边界",
        "已优先使用 support.coinmarketcap.com 官方帮助链接；本地终端直连该页面时遇到 Cloudflare 挑战页，因此正文依据官方检索结果、用户给定更新时间、CMC Academy 官方操作指南和 CMC Disclaimer 中可访问的官方信息整理。",
        "FFF4D6",
        "5C3B00",
    )

    doc.add_heading("1. 这篇官方帮助文章到底在说什么", level=1)
    para(
        doc,
        "文章回答的是一个非常实际的问题：我在 CoinMarketCap 上看到某个币以后，能不能直接买？官方给出的答案不是简单的“能”或“不能”，而是要分链、分入口、分责任。对新手来说，最重要的是先把 CMC 的角色和真正交易发生的位置分清楚。",
    )
    add_table(
        doc,
        ["问题", "小白版答案", "为什么重要"],
        [
            ("CMC 是交易所吗？", "不是。CMC 主要是行情、数据和入口界面。", "不要把 CMC 当成 Binance/Coinbase 那样的账户余额系统。"),
            ("CMC 是钱包吗？", "不是。CMC 不保管私钥，也不保存你的币。", "签名、授权、资产控制权都在你自己的钱包里。"),
            ("可以直接买哪些？", "官方说明当前可通过 DEX Mode 交易 Solana 与 BNB Smart Chain 上的链上代币。", "不是所有 CMC 页面上的币都能在 CMC 内直接交易。"),
            ("其他链怎么办？", "可通过币种详情页的 Trade 或 Markets 区域跳转到第三方 DEX/交易所。", "离开 CMC 后就是第三方服务环境，要重新做安全检查。"),
        ],
        [1.25, 2.45, 2.8],
        8.0,
    )

    doc.add_heading("2. 用 CMC 找买币入口的实际路径", level=1)
    doc.add_heading("2.1 如果是 Solana 或 BSC 链上代币", level=2)
    for item in [
        "打开目标币种的 CoinMarketCap 详情页，确认它所在的网络是否为 Solana 或 BNB Smart Chain。",
        "在页面右上方打开 DEX Mode，页面会出现交易面板。",
        "连接受支持的钱包，例如 Binance Wallet、Trust Wallet、MetaMask、Phantom，或兼容 WalletConnect 的钱包。",
        "选择 Buy 或 Sell，输入数量，查看预计收到数量、路由、滑点、Gas 和服务费信息。",
        "在钱包弹窗里再次核对网络、代币、数量、合约地址和费用，再决定是否签名确认。",
    ]:
        number(doc, item)

    doc.add_heading("2.2 如果不是 Solana 或 BSC 链上代币", level=2)
    for item in [
        "在币种详情页寻找 Trade 按钮；官方说明这可能会把你带到第三方去继续交易。",
        "在 Markets 区域查看哪些中心化交易所或市场交易对列出了该资产。",
        "进入任何第三方页面前，先核对域名、交易对、链和合约地址，不要只相信搜索广告或社群链接。",
        "如果交易所需要注册、KYC、入金或提现，风险就从 CMC 页面转移到该交易所/钱包/链上协议。",
    ]:
        bullet(doc, item)

    add_callout(
        doc,
        "新手误区",
        "在 CMC 上看到一个币、看到一个市场、看到一个 Trade 按钮，都不等于“官方推荐购买”。CMC 官方免责声明强调，网站内容不是投资、金融或交易建议，也不建议用户买入、卖出或持有任何特定加密货币。",
        "FCE8E6",
        "9B1C1C",
    )

    doc.add_heading("3. DEX Mode、第三方服务和钱包风险", level=1)
    para(
        doc,
        "DEX Mode 的关键是“自托管 + 链上执行”：你的资产仍在自己的钱包里，CMC 提供界面和市场数据，DexScan/路由系统帮助寻找第三方 DEX 的价格与流动性。听起来方便，但这也意味着每一笔交易都继承了链上交易的不可逆风险。",
    )
    add_table(
        doc,
        ["风险点", "发生在哪里", "小白应该怎么理解"],
        [
            ("第三方 DEX / AMM", "实际 swap 的流动性池或协议", "CMC 不拥有、运营或审计所有底层协议；池子可能流动性差、价格被操纵或合约有问题。"),
            ("钱包连接", "浏览器钱包或移动钱包", "连接不等于转账，但错误签名、恶意授权或假页面可能造成资产损失。"),
            ("Gas 与协议费", "BSC/Solana 网络与 AMM 协议", "CMC 官方称自身服务费为 0，但你仍可能支付网络 Gas、AMM 费用或代币内置税。"),
            ("高风险代币", "代币合约和交易路径", "官方会筛查疑似 honeypot、rug pull、假币等高风险标记，但筛查不能替代你的尽调。"),
            ("链上不可逆", "钱包确认后写入链上", "一旦交易确认，通常不能撤销、退款或让客服帮你恢复。"),
        ],
        [1.25, 1.95, 3.3],
        7.7,
    )

    doc.add_heading("4. 下单前必须懂的五个关键词", level=1)
    add_table(
        doc,
        ["关键词", "通俗解释", "检查动作"],
        [
            ("滑点", "你看到的价格和最终成交价格之间可能不同。波动大、池子浅、交易大时更明显。", "先小额测试；看 slippage 设置；不要在剧烈波动时用过高容忍度。"),
            ("路由", "系统可能经过一个或多个池子帮你换到目标币。", "确认最终收到的币、数量、网络和路径，不只看按钮上的 Buy。"),
            ("授权", "BSC 等 EVM 链上，首次交易某代币可能要先授权路由合约动用该代币。", "只给必要额度；定期用钱包或授权管理工具撤销不再需要的授权。"),
            ("假币", "名字、图标、ticker 都可能被仿冒，真正识别对象是链和合约地址。", "从项目官网、CMC 页面和区块浏览器交叉核对合约地址。"),
            ("钓鱼", "假 CMC、假钱包弹窗、假 DEX、假客服都可能诱导你签名或输入助记词。", "手输或收藏官方域名；任何页面都不应索要助记词或私钥。"),
        ],
        [0.95, 2.85, 2.7],
        7.7,
    )

    doc.add_heading("5. 安全检查清单", level=1)
    para(doc, "这部分可以直接当作每次买币前的核对清单。新手最容易亏钱的地方，往往不是不知道怎么买，而是没有在签名前慢下来检查。")
    checks = [
        "确认你访问的是 coinmarketcap.com 或 support.coinmarketcap.com 的官方域名，不从广告、群聊短链进入。",
        "确认目标资产的链、合约地址、ticker、图标和项目官网一致；同名币和仿盘很多。",
        "确认钱包网络正确，BSC 用 BNB 做 Gas，Solana 用 SOL 做 Gas；钱包里预留足够 Gas。",
        "确认预计收到数量、最低收到数量、滑点容忍度、Gas、协议费和可能的代币税。",
        "首次授权时只给必要权限；大额资产不要长期暴露在无限授权下。",
        "不要用主钱包做第一次测试；先用小额、干净的钱包试一次完整流程。",
        "看到高收益、空投、客服私信、恢复资产、代操作等说法，默认按钓鱼处理。",
        "交易完成后在区块浏览器或钱包历史里确认交易哈希、成交数量和剩余余额。",
        "如果交易失败，不要连续盲点确认；先查失败原因和是否已经扣除 Gas。",
        "买入前写下理由、最大可亏金额和退出条件；不要把 CMC 排名或热度当成买入理由。",
    ]
    for item in checks:
        bullet(doc, item)

    doc.add_page_break()
    doc.add_heading("6. 给新手的推荐操作顺序", level=1)
    add_table(
        doc,
        ["阶段", "应该做", "不要做"],
        [
            ("学习阶段", "先读 CMC 帮助文章、Academy 指南、免责声明，理解 DEX Mode 和自托管。", "不要看到教程就立刻大额买入。"),
            ("验证阶段", "用小额钱包在官方页面走一遍连接、报价、签名、区块浏览器检查。", "不要连接主钱包，不要授予不必要的无限额度。"),
            ("交易阶段", "只买自己研究过、能承受亏损的金额；记录价格、理由、哈希和费用。", "不要追涨、听群友喊单、忽视滑点和代币税。"),
            ("复盘阶段", "检查实际成交价、费用、滑点、路由和钱包授权，必要时撤销授权。", "不要把一次成功交易误认为流程永远安全。"),
        ],
        [1.05, 3.0, 2.45],
        7.8,
    )

    add_callout(
        doc,
        "最终提醒",
        "CoinMarketCap 的价值是信息聚合和入口便利，不是风险兜底。你点击 Trade、连接钱包、授权合约、签名交易时，每一步都应当当作真实资金操作处理。",
        "E8F4FD",
        "0B2545",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("7. 官方来源与引用边界", level=1)
    para(
        doc,
        "本总结优先围绕 CoinMarketCap 官方帮助文章展开，并用 CMC Academy 操作指南和 CMC Disclaimer 补充解释。由于本地自动抓取官方帮助页时触发 Cloudflare 挑战页，文档没有声称逐字复刻原文，而是对官方可访问信息进行中文归纳。",
        size=9.6,
    )
    for label, url in SOURCES:
        source_line(doc, label, url)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
