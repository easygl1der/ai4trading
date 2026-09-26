from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/yitwah/Documents/Codex/2026-06-07/video-youtube-video-open-document-recently/outputs/crypto_algorithmic_trading_mastery_summary_zh.docx")


def rgb(value):
    return RGBColor.from_string(value)


def set_font(run, size=10, color="000000", bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def cell_text(cell, text, bold=False, size=8.2, color="000000", fill=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            set_cell_width(row.cells[i], width)


def add_table(doc, headers, rows, widths, size=7.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell_text(table.rows[0].cells[i], header, bold=True, size=8.3, color="0B2545", fill="E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, size=size)
    set_table_widths(table, widths)
    doc.add_paragraph("")
    return table


def add_callout(doc, title, body, fill="FFF4D6", color="5C3B00"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(f"{title}：")
    set_font(r, size=9.2, color=color, bold=True)
    r2 = p.add_run(body)
    set_font(r2, size=9.2, color=color)
    doc.add_paragraph("")


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=9.6)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=9.6)


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_font(r, size=10, bold=True)
        text = text.removeprefix(bold_prefix)
    r = p.add_run(text)
    set_font(r, size=10)


def source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{label}: {url}")
    set_font(r, size=8.4, color="555555")


def setup_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, bold, before, after in [
        ("Title", 23, "0B2545", False, 0, 8),
        ("Heading 1", 15.5, "0B2545", True, 13, 6),
        ("Heading 2", 12.5, "1F4D78", True, 10, 5),
        ("Heading 3", 11, "1F4D78", True, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build():
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("《Crypto Algorithmic Trading Mastery》中文导读与实操总结")

    sub = doc.add_paragraph()
    r = sub.add_run("作者：Khushabu Gupta｜平台：Google Play Books｜出版时间：2025-10｜篇幅：约 62 页｜整理日期：2026-06-07")
    set_font(r, size=9.5, color="555555")

    add_callout(
        doc,
        "资料边界",
        "我没有绕过 Google Play/Google Books 的购买、登录或 DRM 机制，也没有获得可合法下载的完整正文文件。本文基于你提供的书籍元数据、Google Play Books 的合法导出说明，以及 Freqtrade、Hummingbot、CCXT、Jesse 等官方文档，整理成中文导读和可执行学习路线。若你已在 Google Play 购买或获取此书，可按 Google 官方帮助从个人书库导出到电脑后再做逐页精读版总结。",
    )

    add_callout(
        doc,
        "一句话结论",
        "这本 62 页小册适合当作“自动化交易入门目录”：知道 Python bot、回测、交易所 API、AI 辅助研究各自是什么。但它篇幅短，不能单独支撑实盘；真正动手时必须配合官方工具文档、交易所 API 文档和严格风控。",
        "E8F4FD",
        "0B2545",
    )

    doc.add_heading("1. 如何合法获取和下载", level=1)
    para(doc, "如果这本书在你的 Google Play Books 账号中可读，合法路径是从个人书库进入书籍页面，然后按 Google Play Help 的说明导出到电脑。部分书籍只允许下载 ACSM/EPUB/PDF 中的一种格式，也可能因为出版社或 DRM 限制无法导出完整无保护文件。")
    for item in [
        "不要从不明网盘、Telegram 群或盗版站下载所谓 PDF；这类文件常见风险包括木马、钓鱼链接、篡改内容和侵权。",
        "如果只能在线阅读，可以用 Google Play 的网页阅读器或移动端离线阅读功能学习；需要做精读笔记时，把你合法获取的摘录或章节标题发给我即可继续整理。",
        "如果你能导出 EPUB/PDF，后续可以做一版“逐章逐页”中文总结、术语表、代码练习和复习题。",
    ]:
        bullet(doc, item)

    doc.add_heading("2. 这本书适合怎么读", level=1)
    add_table(
        doc,
        ["读者阶段", "应该抓住什么", "不要误解成什么"],
        [
            ("完全小白", "理解 bot、回测、API、策略、AI 辅助这些词的基本含义。", "不要把短书当成稳赚交易系统。"),
            ("会一点 Python", "把书里的概念映射到 pandas、CCXT、Freqtrade/Jesse 回测流程。", "不要直接复制任何策略上实盘。"),
            ("想做自动化交易", "关注系统结构：数据、信号、风控、执行、日志、监控、停机。", "不要只盯着“AI”“套利”“收益率”。"),
            ("想做 AI 交易", "把 AI 当研究助手、特征工程和文本/情绪分析工具。", "不要让模型直接控制资金和下单权限。"),
        ],
        [1.15, 3.1, 3.25],
        8.0,
    )

    doc.add_heading("3. 章节式中文总结", level=1)
    doc.add_heading("3.1 Crypto 自动化交易的核心思想", level=2)
    para(doc, "算法交易不是“机器自动赚钱”，而是把一套明确规则交给程序持续执行。它的优势是速度、纪律、可重复和可记录；它的弱点是规则可能失效、市场状态会变化、交易所和网络会出问题。")
    bullet(doc, "人工交易常见问题：情绪化、追涨杀跌、忘记止损、复盘不完整。")
    bullet(doc, "自动化交易常见问题：回测过拟合、忽略手续费/滑点、异常订单、API 限速、策略在行情切换后失效。")
    bullet(doc, "新手最先要学的是系统和风险，不是寻找“神奇指标”。")

    doc.add_heading("3.2 Python 交易 bot 的基本结构", level=2)
    add_table(
        doc,
        ["模块", "作用", "新手要检查的坑"],
        [
            ("数据层", "从交易所取 ticker、OHLCV、订单簿、成交和账户数据。", "时间戳、缺失 K 线、API 限速、不同交易所字段含义不同。"),
            ("信号层", "把数据转换成买入、卖出、观望、调仓等信号。", "未来函数、过度调参、只在牛市有效。"),
            ("风控层", "限制单笔风险、总仓位、最大回撤、止损、冷却期。", "没有风控的 bot 不是交易系统。"),
            ("执行层", "下单、撤单、查询订单状态、处理失败和重复请求。", "市价单滑点、订单半成交、网络超时后重复下单。"),
            ("记录层", "保存行情快照、信号、订单、成交、错误和资产曲线。", "没有日志就无法复盘，也无法定位亏损原因。"),
            ("监控层", "发现异常、发送通知、自动停机或降级。", "只会运行不会停机，是实盘大风险。"),
        ],
        [0.85, 3.0, 3.35],
        7.8,
    )

    doc.add_heading("3.3 回测不是证明赚钱，而是排除明显错误", level=2)
    para(doc, "回测的作用是回答“如果当时按这些规则执行，会发生什么”。它不能证明未来收益，只能帮助你发现策略是否逻辑自洽、是否承受得住回撤、是否对手续费和滑点过于敏感。")
    for item in [
        "必须加入手续费、滑点和最小下单量；crypto 交易成本会明显改变策略结果。",
        "必须避免 look-ahead bias：不能用未来 K 线的信息决定过去的交易。",
        "必须做样本外测试：只在训练区间好看的策略，很可能是过拟合。",
        "必须看最大回撤、亏损持续时间、交易次数、胜率、盈亏比，而不是只看总收益。",
        "必须从 dry-run / paper trading 过渡，观察真实行情下的延迟、挂单、滑点和订单失败。",
    ]:
        bullet(doc, item)

    doc.add_page_break()
    doc.add_heading("3.4 常见策略类型", level=2)
    add_table(
        doc,
        ["策略", "基本想法", "适合工具", "主要风险"],
        [
            ("趋势跟随", "涨势中买入，跌势中退出或做空。", "Freqtrade / Jesse", "震荡行情反复止损。"),
            ("均值回归", "价格偏离均值后赌回归。", "Freqtrade / Jesse", "极端行情中越跌越买。"),
            ("突破策略", "价格突破区间后跟随动量。", "Freqtrade / Jesse", "假突破、滑点。"),
            ("网格", "在区间内低买高卖，反复挂单。", "Hummingbot / 交易所 bot", "单边行情套牢，手续费侵蚀。"),
            ("做市", "同时挂买卖单赚 spread。", "Hummingbot", "库存偏移、撤单延迟、被单边行情打穿。"),
            ("跨所套利", "不同交易所价格差超过成本时搬运或对冲。", "CCXT / Hummingbot", "提现延迟、账户资金分散、风控复杂。"),
            ("资金费率套利", "现货/永续或合约间对冲，赚 funding 差。", "CCXT + 自建系统", "基差变化、爆仓、资金费率反转。"),
        ],
        [1.05, 2.25, 1.65, 2.6],
        7.7,
    )

    doc.add_heading("3.5 AI 在交易中的正确位置", level=2)
    para(doc, "AI 更适合当研究和工程助手，而不是自动替你管理资金。它可以帮助整理交易日志、解释指标、生成初版代码、做新闻/社媒情绪分类、提出假设和发现异常，但不应该直接拥有提现权限或未经人工审查就实盘下单。")
    bullet(doc, "适合 AI 的任务：代码模板、日志摘要、策略假设生成、研究报告摘要、风险清单、异常订单解释。")
    bullet(doc, "不适合 AI 直接做的任务：决定重仓、开高杠杆、绕过风控、读取私钥、控制提现、把社媒消息直接转成买卖。")
    bullet(doc, "AI 交易 bot 的核心风险是幻觉和不可复现：模型说得很像对，但不代表策略在市场中有效。")

    doc.add_heading("4. 工具映射：读完这本书后该用什么", level=1)
    add_table(
        doc,
        ["目标", "首选工具", "为什么", "第一件事"],
        [
            ("学习自动交易系统", "Freqtrade", "有策略、回测、dry-run、FreqUI/REST/Telegram，适合小白从规则策略开始。", "跑官方 Docker quickstart，再写一个极简 RSI/均线策略。"),
            ("学习交易所 API", "CCXT", "统一封装多交易所 public/private API，适合学行情、订单簿、余额、下单。", "先只抓公开 OHLCV 和 ticker，不配置私钥。"),
            ("做回测研究和可视化", "Jesse", "重视回测、图表、导出、优化和研究流程。", "用官方 Docker 启动，跑一个历史回测，看 drawdown 和 trades。"),
            ("做市/网格/跨所执行", "Hummingbot", "面向 market making、controllers、executors 和长期运行策略。", "先模拟 simple PMM 或 grid，理解 inventory 和 spread。"),
        ],
        [1.45, 1.35, 3.05, 2.15],
        7.8,
    )

    doc.add_heading("5. 30 天入门路线", level=1)
    for step in [
        "第 1-3 天：只学 CEX、DEX、现货、永续、限价单、市价单、maker/taker、手续费、滑点。",
        "第 4-7 天：用 CCXT 只抓公开行情，不使用 API key；保存 OHLCV 到 CSV，画一张价格和成交量图。",
        "第 8-12 天：安装 Freqtrade，下载历史数据，跑官方示例策略回测；只看指标，不改参数追求高收益。",
        "第 13-16 天：写一个最简单策略，加入止损、最大仓位、交易次数限制；记录每次交易理由。",
        "第 17-20 天：做 dry-run，不真实下单；比较回测表现和实时模拟表现的差异。",
        "第 21-24 天：用 Jesse 或 Freqtrade 输出交易列表，检查最大亏损交易、连续亏损、手续费占比。",
        "第 25-27 天：阅读 Hummingbot 的做市文档，只模拟 spread、inventory、order refresh，不碰实盘。",
        "第 28-30 天：写实盘前检查清单；如果还不能解释每个风险项，就继续模拟，不上实盘。",
    ]:
        number(doc, step)

    doc.add_heading("6. 实盘前最低安全清单", level=1)
    for item in [
        "API key 只开交易权限，不开提现；为 bot 单独创建子账户或独立 key。",
        "每个策略有最大单笔亏损、单日亏损、总仓位上限、最大同时订单数。",
        "交易所异常、网络中断、订单状态未知、余额异常时，bot 能停机而不是继续加仓。",
        "所有订单、信号、错误、余额变化都有日志；日志能回放一次事故。",
        "回测、样本外测试、dry-run、小资金实盘四步都完成，不跳级。",
        "不使用高杠杆，不因为短期盈利加大仓位，不把网格/做市当作稳定收益。",
    ]:
        bullet(doc, item)

    doc.add_heading("7. 推荐配套官方资料", level=1)
    add_table(
        doc,
        ["资料", "用途", "链接"],
        [
            ("Google Play Help", "说明如何从 Google Play Books 将电子书导出到电脑。", "https://support.google.com/googleplay/answer/179863?hl=en"),
            ("Freqtrade Docs", "自动化交易、策略、回测、dry-run、部署。", "https://www.freqtrade.io/en/stable/"),
            ("CCXT Manual", "交易所统一 API、行情、订单、认证、rate limit。", "https://docs.ccxt.com/"),
            ("Hummingbot Docs", "做市、V2 controllers/executors、连接交易所。", "https://hummingbot.org/"),
            ("Jesse Docs", "回测、策略研究、优化、可视化、live trading。", "https://docs.jesse.trade/"),
        ],
        [1.45, 2.0, 4.25],
        7.4,
    )

    doc.add_heading("8. 参考链接", level=1)
    source(doc, "Google Play Books 合法导出说明", "https://support.google.com/googleplay/answer/179863?hl=en")
    source(doc, "Google Play Books 书籍入口（可能需要登录/地区/购买权限）", "https://play.google.com/store/books/details?id=dGOLEQAAQBAJ")
    source(doc, "Freqtrade 官方文档", "https://www.freqtrade.io/en/stable/")
    source(doc, "CCXT 官方文档", "https://docs.ccxt.com/")
    source(doc, "Hummingbot 官方文档", "https://hummingbot.org/")
    source(doc, "Jesse 官方文档", "https://docs.jesse.trade/")

    add_callout(
        doc,
        "最后建议",
        "把这本书当成 1 小时读完的自动化交易地图。真正学习时，先用 Freqtrade 和 CCXT 建立基本功；等能解释回测偏差、手续费、滑点、订单失败和最大回撤后，再研究 Hummingbot 做市或更复杂的 AI/套利系统。",
        "E8F4FD",
        "0B2545",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
