from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/yitwah/Documents/Codex/2026-06-07/video-youtube-video-open-document-recently")
OUT = ROOT / "outputs/youtube_v8_crypto_trading_bot_masterclass_notes_zh.docx"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_font(run, size=10.0, color="000000", bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, bottom=90, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches: float):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            set_cell_width(row.cells[idx], width)
            set_cell_margins(row.cells[idx])


def cell_text(cell, text, *, bold=False, size=8.0, color="000000", fill=None, align=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    set_font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, *, size=7.6, header_fill="E8EEF5"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header in enumerate(headers):
        cell_text(
            table.rows[0].cells[idx],
            header,
            bold=True,
            size=8.1,
            color="0B2545",
            fill=header_fill,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cells[idx], value, size=size, align=align)
    set_table_widths(table, widths)
    doc.add_paragraph("")
    return table


def add_callout(doc, title, body, *, fill="FFF4D6", color="5C3B00"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 7.0)
    set_cell_margins(cell, top=130, bottom=130, start=160, end=160)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.13
    run = p.add_run(f"{title}：")
    set_font(run, size=9.0, color=color, bold=True)
    run = p.add_run(body)
    set_font(run, size=9.0, color=color)
    doc.add_paragraph("")


def para(doc, text, *, size=10.0, color="000000", after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.13
    run = p.add_run(text)
    set_font(run, size=size, color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_font(run, size=9.5)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_font(run, size=9.5)


def source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.06
    run = p.add_run(f"{label}: {url}")
    set_font(run, size=8.0, color="555555")


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13

    for name, size, color, bold, before, after in [
        ("Title", 22.5, "0B2545", False, 0, 8),
        ("Subtitle", 9.4, "555555", False, 0, 8),
        ("Heading 1", 15.5, "0B2545", True, 13, 6),
        ("Heading 2", 12.4, "1F4D78", True, 10, 5),
        ("Heading 3", 10.8, "1F4D78", True, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(9.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.12


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("YouTube V8｜Crypto Trading Bot Masterclass (2026) 中文记录")

    sub = doc.add_paragraph(style="Subtitle")
    sub.add_run(
        "视频：Crypto Tips｜时长：PT13M10S｜发布：2026-03-01｜整理：2026-06-07｜定位：低权重视频，按风险教育材料处理"
    )

    add_callout(
        doc,
        "资料边界",
        "已通过 BrowserOS YouTube MCP 成功读取视频详情和完整英文 transcript。本文不是逐字翻译，也不构成投资建议；视频中的平台、宏观叙事和收益相关表述只作为待验证信息记录，安全与监管风险部分同时参考官方/安全机构资料。",
        fill="E8F4FD",
        color="0B2545",
    )
    add_callout(
        doc,
        "一句话结论",
        "这支视频表面是 2026 crypto trading bot 入门课，实质更像一份“别把 bot 当自动赚钱机器”的风险提醒。对小白有价值的是安全框架、托管/自托管区别、API key 风险和几类 bot 的基本概念；不应把它当作实盘教程或产品推荐清单。",
        fill="FFF4D6",
        color="5C3B00",
    )

    doc.add_heading("1. 视频主旨：学习技能，不是买收益承诺", level=1)
    para(doc, "视频开头先强调多数交易者会亏损，并指出使用 bot 的零售交易者并不会天然优于手动交易者。它给出的核心论点是：API 集成、自动化、策略逻辑和风控是可迁移技能，但 bot 本身不是“被动收入机器”。")
    para(doc, "因此，这份记录把视频标为低权重：可以学习概念和风险识别，但不要采信“睡觉赚钱”“稳赚套利”“无代码快速部署”等宣传语，更不要据此开 API 权限、买付费 bot 或向陌生智能合约转账。")

    doc.add_heading("2. 视频提到的 bot 类型", level=1)
    add_table(
        doc,
        ["层级/类型", "视频中的说法", "小白应如何理解", "主要风险"],
        [
            ("零代码平台", "Pionex、Coinrule、Cryptohopper 等可快速创建规则或内置 bot。", "适合认识网格、DCA、规则触发等概念，不等于适合重仓。", "托管风险、平台安全、费用磨损、API key 暴露。"),
            ("AI 辅助生成", "用 Cursor、Replit、Claude 等把自然语言转成 Python bot。", "AI 可帮你学习代码结构，但生成结果必须人工审查和模拟测试。", "幻觉代码、错误订单逻辑、缺少风控、泄露密钥。"),
            ("本地 AI agent", "视频提到 open-source/local-first agent，强调密钥和数据留在本机。", "本地化能减少云平台信任面，但部署、日志、隔离和安全责任都转移给你。", "配置错误、模型误操作、服务器暴露、依赖供应链风险。"),
            ("自托管开源", "Freqtrade 这类开源 bot 可审计、可回测、可本地运行。", "更适合认真学习，因为能看代码、回测、dry-run、逐步加风控。", "过拟合、回测失真、实盘滑点、交易所 API 与风控失误。"),
        ],
        [1.05, 2.05, 2.45, 2.25],
        size=7.4,
    )

    doc.add_heading("3. 网格、DCA、套利、AI bot 的基本区别", level=1)
    add_table(
        doc,
        ["类型", "基本逻辑", "适合理解的场景", "不是稳赚的原因"],
        [
            ("Grid 网格", "在一段价格区间内分层挂买单和卖单，赚震荡中的价差。", "横盘、波动较高且区间相对稳定的市场。", "单边下跌会越买越套，单边上涨会过早卖飞；手续费会吃掉小价差。"),
            ("DCA 定投/分批", "按时间或价格下跌分批买入，降低一次性买入的时点风险。", "长期学习和仓位纪律，尤其是不想一次性满仓的人。", "资产本身若持续走弱，DCA 只是把亏损摊开；加仓规则可能导致越亏越重。"),
            ("Arbitrage 套利", "利用不同市场、交易对或链上/链下价格差。", "理解价格发现、流动性和跨市场摩擦。", "真实套利受延迟、滑点、手续费、提现限制、风控冻结和 MEV 影响。"),
            ("AI bot", "用 AI 生成代码、解释信号或辅助策略研究。", "学习编程、日志分析、策略解释和回测复盘。", "AI 不知道未来价格；生成代码可能错误，模型输出不能替代风控和审计。"),
        ],
        [0.95, 2.35, 2.15, 2.35],
        size=7.45,
    )

    doc.add_heading("4. 常见宣传话术与反面样本", level=1)
    para(doc, "视频多次提醒“passive income while you sleep”“risk-free profits”“AI bot narratives”等话术的危险性。本文把这些话术当作反面样本：它们利用小白对通胀、储蓄收益低和错失机会的焦虑，把复杂风险包装成简单按钮。")
    for item in [
        "“保证收益”“每天/每月固定收益”：所有交易都有风险，固定收益承诺通常是骗局高危信号。",
        "“无代码、15 分钟上线、无需经验”：可作为学习入口，但绝不能跳过权限、资金、日志和模拟测试。",
        "“部署这个智能合约就能套利”：小白看不懂合约时，实际可能是在授权 drainer 合约转走钱包资产。",
        "“把 API key 接到平台即可被动收入”：即便不开提现权限，也可能被恶意交易、刷量、对敲或错误策略造成损失。",
        "“AI 生成，所以更聪明”：AI 可以写错代码，也可能忽略交易所限制、订单状态、异常重试和资金安全。",
    ]:
        bullet(doc, item)

    doc.add_heading("5. 被动收入/稳赚风险：为什么小白最容易被打中", level=1)
    add_callout(
        doc,
        "风险教育判断",
        "视频把“储蓄跑不赢通胀、系统对储户不友好”的宏观叙事作为引入，这能解释为什么人们想相信 bot，但不能证明任何 bot 能赚钱。越是能击中焦虑的叙事，越要把它和实盘能力、风险披露、第三方审计分开看。",
        fill="FDEBD0",
        color="6B3D00",
    )
    para(doc, "对 crypto 小白来说，真正的危险不是“不懂高级量化”，而是把未知风险误认为自动化收益。交易 bot 只会更快、更机械地执行规则；如果规则本身没有优势，或者执行层有漏洞，自动化会把亏损放大。")

    doc.add_heading("6. 实盘前安全检查清单", level=1)
    for step in [
        "只用小资金学习，先把最大可承受损失写下来；不要动生活费、学费、房租和借来的钱。",
        "API key 只开必要权限；新手严禁开启提现权限；能设 IP 白名单、有效期和子账户就尽量设置。",
        "先跑 backtest，再跑 dry-run/纸面交易，再小资金实盘；每一步都要保存订单、余额和错误日志。",
        "检查策略是否包含止损、最大仓位、最大单日亏损、异常停机、重复下单保护和交易所断线处理。",
        "不要部署看不懂的智能合约；不要复制 YouTube 描述区、Telegram、Discord 里的合约代码直接运行。",
        "审查费用、滑点和流动性：小利润策略如果没把手续费和成交概率算进去，回测再漂亮也没意义。",
        "确认托管模型：资金是在交易所、bot 平台、DEX 钱包还是本地服务里；每一种都对应不同风险。",
        "定期轮换和撤销 API key；发现异常交易、余额变化或日志缺口时，第一动作是停 bot 和撤权限。",
    ]:
        number(doc, step)

    doc.add_heading("7. 对本视频的使用建议", level=1)
    add_table(
        doc,
        ["用途", "是否适合", "原因"],
        [
            ("小白建立风险意识", "适合", "视频把 drainer、API 泄露、托管风险、保证收益话术讲得比较集中。"),
            ("选择平台或产品", "低权重", "描述区含多个商业链接，平台提及不等于客观评测或适配个人情况。"),
            ("直接照做实盘", "不适合", "视频没有给出可验证策略、回测、风控参数和实盘记录。"),
            ("学习自动化交易路线", "部分适合", "可以把零代码、AI 辅助、本地 agent、自托管开源理解为学习阶梯。"),
            ("判断宏观货币观点", "不适合", "视频后半段有明显价值观/宏观叙事，应另用严肃经济资料交叉验证。"),
        ],
        [1.4, 1.05, 4.0],
        size=7.6,
    )

    doc.add_heading("8. 新手最终行动路线", level=1)
    for item in [
        "第一步只看懂概念：网格、DCA、套利、AI 辅助、API key、托管与自托管。",
        "第二步只做模拟：用公开行情数据和纸面账户理解订单、手续费和日志。",
        "第三步学习开源工具：优先看 Freqtrade 这类可回测、可 dry-run、代码可审计的框架。",
        "第四步才考虑极小资金实盘：严格限制仓位、亏损、权限和运行时长。",
        "第五步复盘而不是幻想：记录每笔交易为什么发生、是否符合预期、哪里和回测不同。",
    ]:
        bullet(doc, item)

    doc.add_heading("9. 来源链接", level=1)
    source(doc, "YouTube 视频", "https://www.youtube.com/watch?v=Er4KvQtZpxw")
    source(doc, "BrowserOS YouTube MCP", "成功读取 video details 与完整 transcript；视频 ID Er4KvQtZpxw")
    source(doc, "CFTC: AI Won't Turn Trading Bots into Money Machines", "https://www.cftc.gov/LearnAndProtect/AdvisoriesAndArticles/AITradingBots.html")
    source(doc, "CFTC/SEC: Fraudulent Digital Asset and Crypto Trading Websites", "https://www.cftc.gov/ConsumerProtection/FraudAwarenessPrevention/CFTCFraudAdvisories/watch_out_for_digital_fraud.html")
    source(doc, "FBI: North Korea Responsible for $1.5 Billion Bybit Hack", "https://www.fbi.gov/investigate/cyber/alerts/psa/north-korea-responsible-for-1-5-billion-bybit-hack")
    source(doc, "3Commas: API data disclosure incident", "https://3commas.io/blog/notice-on-api-data-disclosure-incident")
    source(doc, "SentinelOne: Ethereum Drainers Pose as Trading Bots", "https://www.sentinelone.com/labs/smart-contract-scams-ethereum-drainers-pose-as-trading-bots-to-steal-crypto/")

    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
