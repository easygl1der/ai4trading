from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
OUT = BASE / "outputs"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "火烈鸟网络游戏代币转Crypto资产与香港稳定币市场可行性分析.docx"
SCREENSHOT = BASE / "assets" / "question-4-screenshot.png"


BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(34, 34, 34)
MUTED = RGBColor(89, 89, 89)
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F3F5F7"
PALE_YELLOW = "FFF7DF"
PALE_RED = "FCEAEA"


def set_run_font(run, size=None, bold=None, color=None, font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_font(paragraph, size=11, color=DARK, bold=False):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_table_borders(table, color="D9E2EC"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_para(doc, text="", style=None, size=11, bold=False, color=DARK, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        size = 16
        color = BLUE
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        size = 13
        color = BLUE
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        size = 12
        color = RGBColor(43, 87, 128)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(item)
        set_run_font(run, size=10.8, color=DARK)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(item)
        set_run_font(run, size=10.8, color=DARK)


def add_callout(doc, title, body, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="C7D8EA")
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_matrix(doc, headers, rows, widths, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    set_table_width(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=10, bold=True, color=DARK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(value)
            set_run_font(r, size=9.5, color=DARK)
            if i == 0:
                r.bold = True
            set_cell_margins(cells[i])
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin, Inches(1))
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.8)
        style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "火烈鸟网络游戏代币转 Crypto 资产分析"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_font(header, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("内部研究文档 | 2026-06-08")
    set_run_font(run, size=9, color=MUTED)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("案例第 4 问研究报告")
    set_run_font(r, size=13, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("火烈鸟网络是否应利用游戏代币进入香港稳定币市场？")
    set_run_font(r, size=24, bold=True, color=DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("从 CS:GO/CS2 饰品经济、Binance 平台币路径、香港虚拟资产监管与游戏资产化可行性出发")
    set_run_font(r, size=12.5, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    set_table_borders(table)
    set_table_width(table, [Inches(1.5), Inches(4.8)])
    rows = [
        ("研究对象", "火烈鸟网络、游戏内代币/道具、香港稳定币市场"),
        ("核心问题", "游戏资产能否转化为 crypto 资产，并支撑类似 Binance 的资本故事"),
        ("结论倾向", "可探索游戏资产数字化；不宜直接自发稳定币或复制 Binance"),
        ("日期", "2026 年 6 月 8 日"),
    ]
    for idx, (k, v) in enumerate(rows):
        for j, text in enumerate((k, v)):
            cell = table.rows[idx].cells[j]
            if j == 0:
                shade_cell(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=10.5, bold=(j == 0), color=DARK)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    add_callout(
        doc,
        "一句话结论",
        "火烈鸟可以把游戏资产数字化作为长期创新方向，但不应把“游戏币直接稳定币化、复制 Binance、靠香港 IPO 暴富”作为主战略。更可行的是：游戏 IP 与道具资产标准化 + 合规托管/交易合作 + 持牌稳定币支付场景。",
        fill=PALE_YELLOW,
    )
    doc.add_page_break()


def build():
    doc = Document()
    setup_styles(doc)
    add_cover(doc)

    add_heading(doc, "一、问题重述与直接结论", 1)
    add_para(
        doc,
        "截图中的第 4 问并不是单纯问“火烈鸟要不要做游戏”，而是在追问一个更激进的战略设想：一个以游戏发行和平台业务为基础的公司，能否把游戏内代币、可穿戴道具或平台积分转化为香港市场认可的 crypto 资产，进而复制 Binance 的财富路径并实现 IPO。",
    )
    add_callout(
        doc,
        "判断框架",
        "这个问题需要分开回答：文化 IP 延伸是可行方向；游戏资产数字化有中长期价值；但游戏币直接变稳定币、直接做交易所、直接复制 Binance，现实可行性很低。",
    )
    add_numbered(
        doc,
        [
            "火烈鸟应该继续沿游戏、动漫、影视、网文、电竞等上下游文化产业延伸，强化内容和用户生态。",
            "火烈鸟可以探索游戏内道具、皮肤、徽章、会员权益等资产的标准化和数字化，但应先从可审计权益和合规试点开始。",
            "火烈鸟不宜直接把内部游戏币包装成香港稳定币，因为稳定币需要真实储备、持牌监管、赎回机制、反洗钱系统和持续审计。",
            "Binance 的平台币故事可以作为启发，但不应作为可复制模板。Binance 的成功依赖早期监管窗口、全球交易流动性、交易所现金流和平台币权益设计。",
        ]
    )

    add_heading(doc, "二、为什么 CS:GO/CS2 饰品经济值得作为参照", 1)
    add_para(
        doc,
        "CS:GO/CS2 饰品市场说明：游戏内虚拟物品并非天然没有资产属性。只要它们具备稀缺性、可识别性、可交易性、用户共识和持续流动性，就可能形成类似收藏品、NFT 或小型金融市场的价格体系。",
    )
    add_matrix(
        doc,
        ["维度", "CS:GO/CS2 饰品的表现", "对火烈鸟的启发"],
        [
            ("稀缺性", "箱子掉落、绝版活动、稀有 Float 和高端皮肤共同制造稀缺。", "游戏资产必须有清晰发行量、稀缺等级和不可随意增发的规则。"),
            ("流动性", "Steam 市场和第三方平台提供持续买卖盘。", "没有合规交易和托管机制，游戏权益难以成为外部资产。"),
            ("共识", "职业选手、主播、玩家社区和交易商共同塑造价格叙事。", "资产化必须依赖真实用户偏好，而不是只靠公司宣传。"),
            ("价格传导", "箱子、皮肤、磨损、系列、平台汇率之间存在联动。", "火烈鸟若做资产化，需要建立统一定价、库存、交易和披露体系。"),
        ],
        [Inches(1.15), Inches(2.65), Inches(2.65)],
    )
    add_para(
        doc,
        "但是，CS2 饰品不是稳定币。它的价格来自稀缺、审美、社群共识、主播效应和投机资金，而不是一比一法币储备。因此，它可以证明“游戏资产有资产化潜力”，不能证明“游戏币可以直接变成合规稳定币”。",
    )

    add_heading(doc, "三、Binance 为什么会被题目拿来类比", 1)
    add_para(
        doc,
        "Binance 的发家路径通常被概括为“交易所 + 平台币”双轮驱动：交易所提供全球加密资产交易撮合、流动性和手续费现金流；平台币 BNB 则把交易所增长、手续费折扣、平台权益和投机预期绑定在一起。",
    )
    add_matrix(
        doc,
        ["比较项", "Binance / BNB", "火烈鸟 / 游戏代币"],
        [
            ("基础业务", "全球 crypto 交易所，核心是撮合、流动性和手续费。", "手游研发、发行、渠道和平台运营，核心是内容和用户。"),
            ("代币属性", "平台币，绑定手续费折扣、平台权益和后续公链生态。", "更接近游戏点券、道具币、会员权益或平台积分。"),
            ("现金流支撑", "交易量越大，手续费和平台币叙事越强。", "游戏币主要用于消费，不天然连接全球金融流动性。"),
            ("监管窗口", "早期加密市场监管不成熟，存在高速扩张窗口。", "当前香港已建立持牌稳定币和虚拟资产交易平台监管框架。"),
        ],
        [Inches(1.25), Inches(2.55), Inches(2.55)],
    )
    add_callout(
        doc,
        "类比的价值",
        "题目提 Binance，不是说火烈鸟可以照抄 Binance，而是要求讨论：内部平台价值能否通过代币外部化被放大，以及这种放大在当前香港监管环境下还剩多少现实空间。",
        fill=PALE_YELLOW,
    )

    add_heading(doc, "四、香港市场：机会真实存在，但监管门槛已经很高", 1)
    add_para(
        doc,
        "香港之所以出现在题目中，是因为它正在建设较完整的虚拟资产监管框架，是内地背景企业最容易想象的合规 crypto 试验地之一。但“香港允许虚拟资产创新”并不等于“任何游戏公司都可以把积分改名为稳定币”。",
    )
    add_bullets(
        doc,
        [
            "香港金管局公布，稳定币发行人监管制度于 2025 年 8 月 1 日生效。",
            "在香港发行法币参考稳定币，重点是持牌、储备资产、赎回安排、风险管理、审计和反洗钱。",
            "香港证监会要求在香港经营或主动向香港投资者推广服务的中心化虚拟资产交易平台取得牌照。",
            "内地背景企业还必须处理用户来源、宣传边界、资金跨境流动、未成年人保护、黑产和数据合规等问题。",
        ]
    )
    add_matrix(
        doc,
        ["监管要求", "稳定币/交易平台语境", "对游戏币路线的影响"],
        [
            ("真实储备", "稳定币需要高质量、高流动性储备资产支持。", "游戏道具价格波动大，不能直接充当稳定币储备。"),
            ("持牌经营", "发行、交易、托管、推广可能分别触发不同许可。", "火烈鸟需要合作持牌机构，不能只靠技术上链。"),
            ("赎回机制", "用户应能按规则赎回法币或等价资产。", "普通游戏积分通常只可消费，不具备稳定赎回承诺。"),
            ("AML/KYC", "需识别用户、监测交易、处理可疑资金流。", "游戏交易一旦金融化，会显著增加合规和运营成本。"),
            ("投资者保护", "需披露风险、限制不适当销售和欺诈宣传。", "不能用“暴富”“复制 Binance”等叙事面向公众推广。"),
        ],
        [Inches(1.25), Inches(2.6), Inches(2.5)],
    )

    add_heading(doc, "五、火烈鸟的现实能力边界", 1)
    add_para(
        doc,
        "从现有材料看，火烈鸟网络更像一家移动游戏研发、发行和渠道平台公司。它可能拥有用户、渠道、发行能力和一定海外覆盖，但并不天然拥有交易所基础设施、金融合规团队、链上风控系统、资产托管能力和稳定币储备管理能力。",
    )
    add_matrix(
        doc,
        ["能力项", "火烈鸟可能已有基础", "缺口"],
        [
            ("内容与用户", "游戏发行、平台渠道、玩家触达、运营活动。", "需要更强 IP 生产能力和长期社区共识。"),
            ("资产设计", "可穿戴道具、游戏币、会员权益、活动奖励。", "需要标准化编号、发行量披露、权利边界和二级流转规则。"),
            ("交易能力", "可能有内部兑换、充值、消费系统。", "缺少持牌交易、托管、KYC/AML、市场监控和纠纷处理体系。"),
            ("金融合规", "传统游戏业务合规经验。", "稳定币发行、虚拟资产平台、跨境支付、储备审计能力不足。"),
        ],
        [Inches(1.25), Inches(2.6), Inches(2.5)],
    )
    add_para(
        doc,
        "因此，火烈鸟最适合扮演的角色不是“稳定币发行人”或“交易所”，而是“游戏内容和资产场景提供方”。真正的金融环节应尽量通过持牌机构完成。",
    )

    add_heading(doc, "六、可行战略：三阶段推进，而不是一步发币", 1)
    add_heading(doc, "阶段一：文化 IP 和游戏资产标准化", 2)
    add_para(
        doc,
        "第一阶段不应急于发币，而应先把游戏资产做成可识别、可审计、可解释的权益体系。这包括道具分类、发行量、稀缺等级、使用期限、转让限制、销毁规则、用户权利和平台权利。",
    )
    add_bullets(
        doc,
        [
            "区分纯消费品、可收藏品、会员权益、赛事纪念品和游戏外权益票券。",
            "建立统一资产编号、库存披露、发行公告和历史记录。",
            "限制随意增发，避免平台信用被破坏。",
            "优先选择国产文化 IP、赛事、联名皮肤和高情绪价值道具。"
        ]
    )
    add_heading(doc, "阶段二：合规数字藏品或托管凭证试点", 2)
    add_para(
        doc,
        "第二阶段可以把少量高价值资产做成链上凭证、托管凭证或数字藏品，但交易、托管和用户身份识别应交给合规合作方。这个阶段的目标是验证玩家是否真的愿意为可收藏、可转让、可展示的权益付费。",
    )
    add_heading(doc, "阶段三：接入持牌稳定币或支付网络", 2)
    add_para(
        doc,
        "第三阶段才考虑稳定币，但火烈鸟更合理的角色是使用持牌稳定币或支付网络，而不是自己发行稳定币。例如，面向海外玩家用持牌稳定币支付游戏服务、做跨境结算、降低海外发行收款摩擦。稳定币在这里是支付工具，不是游戏币本身。",
    )

    add_heading(doc, "七、风险清单", 1)
    add_matrix(
        doc,
        ["风险", "具体表现", "建议处理"],
        [
            ("监管风险", "游戏积分被宣传为投资品、稳定币或可升值资产。", "避免收益承诺；所有金融功能通过持牌机构完成。"),
            ("资产泡沫", "玩家因投机而非使用需求买入道具，价格崩盘后损害品牌。", "限制杠杆、限制未成年人交易、提高披露透明度。"),
            ("平台信用", "平台随意增发、改规则或封禁账号导致资产信任崩塌。", "建立公开规则、变更公告和用户申诉机制。"),
            ("黑产洗钱", "低门槛道具交易被刷量、盗号、洗钱利用。", "接入 KYC、交易监控、异常账户冻结和黑名单。"),
            ("战略跑偏", "从游戏公司变成高风险金融公司，丢失主业优势。", "坚持内容和场景为主，金融能力通过合作获得。"),
        ],
        [Inches(1.25), Inches(2.55), Inches(2.55)],
        header_fill=PALE_RED,
    )

    add_heading(doc, "八、可以直接用于作业的完整回答", 1)
    paragraphs = [
        "火烈鸟网络应当继续向上下游文化产业延伸，尤其是国产游戏、动画、影视、网文、电竞和虚拟道具等方向。国内爆款文化产品说明，中国年轻用户愿意为优质内容、身份表达和情绪价值付费，这与火烈鸟的游戏发行和平台运营能力存在协同。相较于单纯做渠道，向 IP、内容和玩家资产体系延伸，更有可能形成长期壁垒。",
        "但是，火烈鸟不宜直接把游戏中积累的可穿戴币、道具币或平台积分包装成香港稳定币。稳定币在香港已经进入持牌监管阶段，发行人需要真实储备、审计、赎回机制、反洗钱系统和本地合规安排。游戏币本质上更接近平台内预付费积分或虚拟道具权益，其价值依赖游戏热度和玩家共识，并不适合作为稳定币本体或储备资产。",
        "CS:GO/CS2 饰品市场说明，游戏资产确实可以因为稀缺性、审美、社群共识和交易流动性而形成价格体系。但这类资产更接近数字收藏品或平台内许可权益，而不是法币锚定稳定币。火烈鸟可以借鉴 CS2 饰品经济中的资产标准化、稀缺发行和社区运营经验，却不能直接把这种投机价格当成稳定币信用基础。",
        "Binance 的故事也只能作为启发，而不是火烈鸟可照抄的模板。Binance 的成功来自早期监管窗口、全球交易流动性、交易所手续费现金流和平台币经济模型。火烈鸟是游戏发行和渠道公司，并不天然具备交易所的撮合能力、金融合规能力和全球流动性基础。如果强行走“游戏币稳定币化 + 自建交易所 + 复制 Binance”的路线，监管成本、合规风险和战略偏离都会非常高。",
        "更合理的方案是：火烈鸟先把游戏内可收藏资产标准化，探索数字藏品、链上凭证和合规托管交易；再与香港持牌虚拟资产平台、托管机构、支付机构或稳定币发行人合作，把自己定位为内容和场景提供方，而不是金融牌照主体。这样既可以抓住游戏资产数字化趋势，也能避免直接承担稳定币发行和交易所运营的高监管风险。",
        "因此，火烈鸟未来的资本故事应当是“游戏发行平台 + 文化 IP 生态 + 合规数字资产创新”，而不是“游戏币稳定币化 + 复制 Binance 暴富”。如果未来要 IPO，也应以稳定的游戏业务收入、IP 运营能力、用户规模和合规创新为主线，把 crypto 相关业务作为谨慎试点和增值场景，而不是核心融资叙事。",
    ]
    for text in paragraphs:
        add_para(doc, text)

    add_heading(doc, "九、截图原题与材料来源", 1)
    add_para(doc, "下图为本报告回答的案例第 4 问截图。")
    if SCREENSHOT.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(SCREENSHOT), width=Inches(6.2))
    add_para(
        doc,
        "本报告综合了此前两份材料：一份关于 CS:GO/CS2 饰品价格机制、火烈鸟与稳定币可行性的完整报告；另一份关于微信回答语气、Binance 发家史与香港市场关系的聊天稿。",
        size=10.5,
        color=MUTED,
    )

    add_heading(doc, "十、参考资料", 1)
    refs = [
        "香港金管局：Implementation of regulatory regime for stablecoin issuers，2025-07-29。https://www.hkma.gov.hk/eng/news-and-media/press-releases/2025/07/20250729-4/",
        "香港证监会：Virtual asset trading platform operators。https://www.sfc.hk/en/Rules-and-standards/Virtual-assets/Virtual-asset-trading-platforms-operators",
        "香港证监会：Circular on transitional arrangements of the new licensing regime for virtual asset trading platforms。https://apps.sfc.hk/edistributionWeb/gateway/EN/circular/doc?refNo=23EC27",
        "本地原始材料：csgo-cs2-skins-and-flamingo-token-stablecoin-analysis.md",
        "本地原始材料：flamingo-question-4-chat-draft-and-binance-context.md",
    ]
    add_bullets(doc, refs)

    doc.save(DOCX)
    return DOCX


if __name__ == "__main__":
    path = build()
    print(path)
